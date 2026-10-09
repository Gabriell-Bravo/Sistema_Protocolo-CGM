"""Smoke tests das telas por papel e regras de HTTP (GET não grava,
POST obrigatório, respostas de erro sem exceção interna)."""

import json

from django.test import override_settings
from django.urls import reverse
from django.utils import timezone

from processos_app.models import EspecieProcesso, EventoProcesso, ProcessHistory, Processo
from processos_app.services import tramitacao

from .base import LIQ, BaseProcessoTestCase, criar_usuario

STATIC = {'STORAGES': {
    'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
}}


@override_settings(**STATIC)
class TelasPorPapelTest(BaseProcessoTestCase):

    def setUp(self):
        self.processo = self.novo_processo()
        self.monitorado = self.novo_processo(self.especie_monit, numero_processo='55/2026')

    def _get(self, usuario, nome, *args, esperado=200, query=''):
        self.client.force_login(usuario)
        resposta = self.client.get(reverse(nome, args=args) + query)
        self.assertEqual(resposta.status_code, esperado,
                         f'{nome} com {usuario.username}: {resposta.status_code}')
        return resposta

    def test_protocolo(self):
        for nome in ('listar_processos', 'listar_finalizados', 'cadastrar_processo'):
            self._get(self.protocolo, nome)
        self._get(self.protocolo, 'processos_para_retirada', esperado=302)
        self._get(self.protocolo, 'ver_historico_processo', self.processo.id)
        self._get(self.protocolo, 'listar_processos', query='?situacao=DISPONIVEL&page=1')

    def test_analista(self):
        self._get(self.analista_lic, 'area_analista')
        self._get(self.analista_lic, 'area_analista', query='?filtro=disponiveis')
        self._get(self.analista_lic, 'analista_processo', self.processo.id)
        self._get(self.analista_lic, 'meus_atendimentos')
        self._get(self.analista_lic, 'listar_finalizados')
        self._get(self.analista_lic, 'gestao_liberados_assinatura')
        self._get(self.analista_lic, 'gestao_diligencias')
        self._get(self.analista_lic, 'controle_analise')
        self._get(self.analista_lic, 'controle_analise', query='?filtro=disponiveis')
        self._get(self.analista_lic, 'controle_relatorio', esperado=403)
        self._get(self.analista_liq, 'controle_analise', esperado=403)

    def test_gestao(self):
        for nome in ('gestao_dashboard', 'gestao_processos', 'gestao_diligencias',
                     'gestao_pessoas', 'gestao_cadastros_inicio',
                     'gestao_liberados_assinatura', 'listar_processos',
                     'listar_finalizados'):
            self._get(self.gestao, nome)
        for slug in ('unidades', 'especies', 'prioridades', 'tipos-indisponibilidade'):
            self._get(self.gestao, 'gestao_cadastros', slug)
        self._get(self.gestao, 'gestao_dashboard', query='?inicio=2026-01-01&fim=2026-12-31')
        # item 2: Gestão consulta a tela do processo sem praticar ato.
        self._get(self.gestao, 'analista_processo', self.processo.id)

    def test_acessos_negados(self):
        self._get(self.analista_lic, 'gestao_cadastros_inicio', esperado=403)
        self._get(self.protocolo, 'gestao_dashboard', esperado=403)
        self._get(self.gestao, 'cadastrar_processo', esperado=403)
        # analista de outro grupo não abre o processo
        self._get(self.analista_liq, 'analista_processo', self.processo.id, esperado=403)
        self._get(self.protocolo, 'gestao_diligencias', esperado=403)
        self._get(self.protocolo, 'gestao_liberados_assinatura', esperado=403)

    def test_analista_consulta_finalizados_assinatura_e_diligencias_sem_editar(self):
        from processos_app.services import pendencias

        finalizado = self.processo_disponivel_retirada()
        tramitacao.registrar_saida([finalizado.id], self.protocolo)
        aguardando = self.processo_em_analise()
        aguardando.numero_processo = '2002/2026'
        aguardando.save(update_fields=['numero_processo'])
        tramitacao.liberar_assinatura(aguardando.id, self.analista_lic)
        liquidacao = self.novo_processo(
            self.especie_liq, numero_processo='2003/2026')
        tramitacao.assumir(liquidacao.id, self.analista_liq)
        liquidacao.refresh_from_db()
        self.preencher_analise(liquidacao, numero_despacho=None)
        tramitacao.liberar_assinatura(liquidacao.id, self.analista_liq)

        pagina = self._get(self.analista_lic, 'listar_finalizados')
        self.assertContains(pagina, 'Finalizados')
        self.assertContains(pagina, 'Fila de análise')
        self.assertNotContains(pagina, 'Exportar Excel')
        self.assertFalse(pagina.context['can_edit'])
        self.assertFalse(pagina.context['can_delete'])
        self.assertFalse(pagina.context['can_concluir_monitoramento'])
        self.assertFalse(pagina.context['can_export'])
        self.assertContains(pagina, finalizado.numero_processo)

        assinar = self._get(self.analista_lic, 'gestao_liberados_assinatura')
        self.assertContains(assinar, '2002/2026')
        self.assertNotContains(assinar, '2003/2026')

        processo_pend = self.processo_em_analise()
        processo_pend.numero_processo = '2004/2026'
        processo_pend.save(update_fields=['numero_processo'])
        pendencia = pendencias.criar(
            processo_pend.id, self.analista_lic, 'Falta documento')
        dilig = self._get(self.analista_lic, 'gestao_diligencias')
        self.assertContains(dilig, 'Falta documento')
        self.assertNotContains(dilig, 'Indicar atendimento')
        self.assertContains(dilig, 'Consulta. A indicação de atendimento é da Gestão.')
        self.assertContains(dilig, 'name="termo"')

        busca = self._get(self.analista_lic, 'gestao_diligencias',
                          query='?termo=2004')
        self.assertContains(busca, 'Falta documento')
        self.assertEqual(busca.context['total'], 1)
        vazia = self._get(self.analista_lic, 'gestao_diligencias',
                          query='?termo=inexistente-xyz')
        self.assertEqual(vazia.context['total'], 0)
        self.assertContains(vazia, 'Nenhuma diligência encontrada')

        self.client.force_login(self.analista_lic)
        resposta = self.client.post(
            reverse('pend_indicar_atendimento', args=[pendencia.id]))
        self.assertEqual(resposta.status_code, 302)
        pendencia.refresh_from_db()
        self.assertEqual(pendencia.status, 'AGUARDANDO_ATENDIMENTO')

    def test_get_nao_grava(self):
        antes = (ProcessHistory.objects.count(), EventoProcesso.objects.count(),
                 list(Processo.objects.values_list('status_monitoramento', flat=True)))
        self._get(self.gestao, 'listar_finalizados')
        self._get(self.gestao, 'listar_processos')
        self._get(self.gestao, 'gestao_dashboard')
        depois = (ProcessHistory.objects.count(), EventoProcesso.objects.count(),
                  list(Processo.objects.values_list('status_monitoramento', flat=True)))
        self.assertEqual(antes, depois)

    def test_gestao_edita_especie_nome_grupo_e_sequencia(self):
        from processos_app.models import EspecieProcesso
        from .base import criar_usuario
        processo = self.novo_processo(self.especie_liq, numero_processo='esp-cad/2026')
        self.assertEqual(processo.genero, 'LIQUIDACOES')
        self.assertEqual(processo.especie, self.especie_liq.nome)

        pagina = self._get(self.gestao, 'gestao_cadastros', 'especies')
        self.assertContains(pagina, 'Espécies de Processo')
        self.assertContains(pagina, 'Grupo do relatório')
        self.assertContains(pagina, 'name="sequencia_numeracao"')

        self.client.force_login(self.gestao)
        resposta = self.client.post(reverse('gestao_cadastro_salvar', args=['especies']), {
            'id': self.especie_liq.id,
            'nome': 'Pagamento Geral Renomeado',
            'grupo': 'LIQUIDACOES',
            'sequencia_numeracao': 'ADIANTAMENTO',
            'ordem': '10',
            'tipo_monitoramento': 'NENHUM',
            'gera_relatorio': 'on',
        })
        self.assertEqual(resposta.status_code, 302)
        self.especie_liq.refresh_from_db()
        self.assertEqual(self.especie_liq.nome, 'Pagamento Geral Renomeado')
        self.assertEqual(self.especie_liq.sequencia_numeracao, 'ADIANTAMENTO')
        processo.refresh_from_db()
        self.assertEqual(processo.especie, 'Pagamento Geral Renomeado')
        self.assertEqual(processo.genero, 'LIQUIDACOES')

        from processos_app.services import relatorios
        self.assertEqual(relatorios.sequencia_do_processo(processo), 'ADIANTAMENTO')

        admin = criar_usuario('admin_esp', 'PROTOCOLO', is_superuser=True)
        self._get(admin, 'gestao_cadastros', 'especies')
        self.client.force_login(admin)
        resposta = self.client.post(reverse('gestao_cadastro_salvar', args=['especies']), {
            'id': self.especie_liq.id,
            'nome': 'Pagamento Geral Admin',
            'grupo': 'LIQUIDACOES',
            'sequencia_numeracao': 'DIARIA',
            'ordem': '10',
            'tipo_monitoramento': 'NENHUM',
            'gera_relatorio': 'on',
        })
        self.assertEqual(resposta.status_code, 302)
        self.especie_liq.refresh_from_db()
        self.assertEqual(self.especie_liq.nome, 'Pagamento Geral Admin')
        self.assertEqual(self.especie_liq.sequencia_numeracao, 'DIARIA')

    def test_exportar_excel(self):
        finalizado = self.processo_disponivel_retirada()
        tramitacao.registrar_saida([finalizado.id], self.protocolo)
        finalizado.refresh_from_db()
        data = finalizado.data_saida.isoformat()

        # Sem período explícito: modelo padrão "hoje" (pode estar vazio).
        hoje = self._get(self.protocolo, 'exportar_finalizados_excel')
        self.assertIn('spreadsheet', hoje['Content-Type'])

        resposta = self._get(
            self.protocolo, 'exportar_finalizados_excel',
            query=f'?periodo=personalizado&data_inicial={data}&data_final={data}')
        self.assertIn('spreadsheet', resposta['Content-Type'])

        todas = self._get(
            self.protocolo, 'exportar_finalizados_excel', query='?periodo=todas')
        self.assertIn('spreadsheet', todas['Content-Type'])

        pagina = self._get(self.protocolo, 'listar_finalizados')
        self.assertContains(pagina, 'Exportar saídas')
        self.assertContains(pagina, reverse('exportar_saidas'))
        self.assertNotContains(pagina, 'modalExportarExcel')

    def test_exportar_ativos_excel_protocolo_e_admin(self):
        from .base import criar_usuario
        from openpyxl import load_workbook
        from io import BytesIO

        ativo = self.novo_processo(numero_processo='exp-ativo/2026')
        pagina = self._get(self.protocolo, 'listar_processos')
        self.assertContains(pagina, 'Gerar planilha')
        self.assertContains(pagina, reverse('exportar_ativos_excel'))

        resp = self._get(self.protocolo, 'exportar_ativos_excel')
        self.assertIn('spreadsheet', resp['Content-Type'])
        wb = load_workbook(BytesIO(resp.content))
        linhas = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(linhas[0][0], 'N° Processo')
        numeros = {row[0] for row in linhas[1:]}
        self.assertIn(ativo.numero_processo, numeros)

        admin = criar_usuario('admin_exp_ativos', 'PROTOCOLO', is_superuser=True)
        resp_admin = self._get(admin, 'exportar_ativos_excel')
        self.assertIn('spreadsheet', resp_admin['Content-Type'])

        # Analista não exporta ativos.
        self.client.force_login(self.analista_liq)
        bloqueado = self.client.get(reverse('exportar_ativos_excel'))
        self.assertEqual(bloqueado.status_code, 403)

    def test_dashboard_mostra_graficos_para_gestao_e_administrador(self):
        from .base import criar_usuario
        self.novo_processo()
        pagina = self._get(self.gestao, 'gestao_dashboard')
        self.assertContains(pagina, 'Onde estão os processos')
        self.assertContains(pagina, 'Prioridade')
        self.assertContains(pagina, 'grafico__barra')
        self.assertGreaterEqual(pagina.context['panorama']['situacao'][0]['valor'], 1)
        admin = criar_usuario('admin_dash', 'PROTOCOLO', is_superuser=True)
        self._get(admin, 'gestao_dashboard')

    def test_administrador_ve_o_botao_de_saida(self):
        from .base import criar_usuario
        self.novo_processo()
        pagina_gestao = self._get(self.gestao, 'listar_processos')
        self.assertNotContains(pagina_gestao, 'Registrar saída')
        admin = criar_usuario('admin_lista', 'GESTAO', is_superuser=True, is_staff=True)
        pagina_admin = self._get(admin, 'listar_processos')
        self.assertContains(pagina_admin, 'Registrar saída')

    def test_gestao_pergunta_se_a_urgencia_e_recorrente(self):
        from processos_app.models import UrgenciaRecorrente
        processo = self.novo_processo(numero_processo='880/2026')
        pagina = self._get(self.gestao, 'gestao_processos')
        self.assertContains(pagina, 'balao-urgencia')
        self.assertContains(pagina, 'só desta vez')
        self.client.post(reverse('gestao_alterar_prioridade', args=[processo.id]), {
            'prioridade': 'URGENTE',
            'urgencia_recorrente': 'sim',
        })
        self.assertTrue(UrgenciaRecorrente.vale_para('880/2026'))
        processo.refresh_from_db()
        self.assertEqual(processo.prioridade, 'URGENTE')

    def test_gestao_e_admin_veem_quem_esta_logado(self):
        pessoas = self._get(self.gestao, 'gestao_pessoas')
        por_nome = {linha['nome']: linha['logado'] for linha in pessoas.context['quadro']}
        self.assertTrue(por_nome['Gestora'])
        self.assertFalse(por_nome['Protocolo'])
        self.assertContains(pessoas, 'Logado')
        self.assertContains(pessoas, 'Não logado')

        admin = self.gestao
        admin.is_superuser = True
        admin.is_staff = True
        admin.save()
        usuarios = self._get(admin, 'manage_users')
        por_id = {u.id: u.esta_logado for u in usuarios.context['users']}
        self.assertTrue(por_id[admin.id])
        self.assertFalse(por_id[self.protocolo.id])
        self.assertContains(usuarios, 'Logado')
        self.assertContains(usuarios, 'Não logado')


@override_settings(**STATIC)
class AcoesHttpTest(BaseProcessoTestCase):

    def _post_json(self, usuario, nome, args, dados):
        self.client.force_login(usuario)
        return self.client.post(reverse(nome, args=args), data=json.dumps(dados),
                                content_type='application/json')

    def test_salvar_processo_json(self):
        resposta = self._post_json(self.protocolo, 'salvar_processo', [],
                                   self.dados_protocolo())
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.assertTrue(resposta.json()['success'])

    def test_gestao_nao_cadastra(self):
        resposta = self._post_json(self.gestao, 'salvar_processo', [], self.dados_protocolo())
        self.assertEqual(resposta.status_code, 403)

    def test_endpoints_de_escrita_recusam_get(self):
        processo = self.novo_processo()
        self.client.force_login(self.gestao)
        for nome in ('deletar_processo', 'marcar_saida_processo', 'concluir_monitoramento',
                     'atualizar_processo'):
            resposta = self.client.get(reverse(nome, args=[processo.id]))
            self.assertEqual(resposta.status_code, 405, nome)

    def test_edicao_em_linha_protocolo_grava_analise_e_saida(self):
        processo = self.novo_processo()
        resposta = self._post_json(self.protocolo, 'atualizar_processo', [processo.id],
                                   {'objeto': 'Alterado', 'numero_despacho': '7/2026',
                                    'data_saida': '2026-10-02', 'hora_saida': '11:00'})
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.assertEqual(resposta.json()['recusados'], [])
        processo.refresh_from_db()
        self.assertEqual(processo.objeto, 'Alterado')
        self.assertEqual(processo.numero_despacho, '7/2026')
        self.assertEqual(processo.data_saida.isoformat(), '2026-10-02')
        self.assertEqual(processo.hora_saida.strftime('%H:%M'), '11:00')

    def test_edicao_em_linha_recusa_campo_de_tramitacao(self):
        processo = self.novo_processo()
        situacao = processo.situacao_tramite
        resposta = self._post_json(self.protocolo, 'atualizar_processo', [processo.id],
                                   {'objeto': 'Ok', 'situacao_tramite': 'SAIDA_CONCLUIDA'})
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.assertIn('situacao_tramite', resposta.json()['recusados'])
        processo.refresh_from_db()
        self.assertEqual(processo.objeto, 'Ok')
        self.assertEqual(processo.situacao_tramite, situacao)

    def test_gestao_nao_edita_em_linha(self):
        processo = self.novo_processo()
        resposta = self._post_json(self.gestao, 'atualizar_processo', [processo.id],
                                   {'objeto': 'Alterado'})
        self.assertEqual(resposta.status_code, 403)

    def test_cancelar_pela_rota_antiga(self):
        processo = self.novo_processo()
        self.client.force_login(self.gestao)
        resposta = self.client.post(reverse('deletar_processo', args=[processo.id]),
                                    {'motivo': 'Duplicado'})
        self.assertEqual(resposta.status_code, 200, resposta.content)
        processo.refresh_from_db()
        self.assertIsNotNone(processo.cancelado_em)

    def test_saida_pelo_botao_funciona_sem_as_etapas_do_meio(self):
        processo = self.novo_processo()
        self.client.force_login(self.protocolo)
        resposta = self.client.post(reverse('marcar_saida_processo', args=[processo.id]))
        self.assertEqual(resposta.status_code, 200, resposta.content)
        processo.refresh_from_db()
        self.assertEqual(processo.situacao_tramite, 'SAIDA_CONCLUIDA')
        self.assertIsNotNone(processo.data_saida)

    def test_gestao_nao_pratica_ato_na_tela_do_processo(self):
        processo = self.novo_processo()
        self.client.force_login(self.gestao)
        resposta = self.client.post(reverse('analista_processo', args=[processo.id]),
                                    {'status_analise': 'NAO_PROSSEGUIMENTO'})
        self.assertEqual(resposta.status_code, 403)

    def test_analista_e_admin_editam_dados_do_resumo(self):
        processo = self.processo_em_analise(self.analista_liq, self.especie_liq)
        self.client.force_login(self.analista_liq)
        tela = self.client.get(reverse('analista_processo', args=[processo.id]))
        self.assertTrue(tela.context['pode_editar'])
        self.assertContains(tela, 'name="objeto"')
        self.assertContains(tela, 'name="contratada"')
        self.assertContains(tela, 'name="volume"')

        resposta = self.client.post(reverse('analista_processo', args=[processo.id]), {
            'objeto': 'Objeto corrigido pelo analista',
            'contratada': 'Empresa XYZ',
            'volume': '5',
            'secretaria': processo.secretaria,
            'especie': processo.especie,
            'genero': processo.genero,
            'prioridade': processo.prioridade or 'NORMAL',
            'data_entrada': processo.data_entrada.isoformat(),
            'hora_entrada': processo.hora_entrada.strftime('%H:%M'),
            'destino': processo.destino or 'Unidade de Teste',
            'valor': '1500',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            'acao': 'salvar',
        })
        self.assertEqual(resposta.status_code, 302, getattr(resposta, 'content', b''))
        processo.refresh_from_db()
        self.assertEqual(processo.objeto, 'Objeto corrigido pelo analista')
        self.assertEqual(processo.contratada, 'Empresa XYZ')
        self.assertEqual(processo.volume, '5')

        admin = criar_usuario('admin_resumo', 'GESTAO', is_superuser=True)
        self.client.force_login(admin)
        tela_admin = self.client.get(reverse('analista_processo', args=[processo.id]))
        self.assertTrue(tela_admin.context['pode_editar'])
        resposta_admin = self.client.post(reverse('analista_processo', args=[processo.id]), {
            'objeto': 'Objeto pelo admin',
            'contratada': 'Empresa XYZ',
            'volume': '5',
            'secretaria': processo.secretaria,
            'especie': processo.especie,
            'genero': processo.genero,
            'prioridade': 'URGENTE',
            'data_entrada': processo.data_entrada.isoformat(),
            'hora_entrada': processo.hora_entrada.strftime('%H:%M'),
            'destino': processo.destino or 'Unidade de Teste',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            'acao': 'salvar',
        })
        self.assertEqual(resposta_admin.status_code, 302)
        processo.refresh_from_db()
        self.assertEqual(processo.objeto, 'Objeto pelo admin')
        self.assertEqual(processo.prioridade, 'URGENTE')

    def test_direcionar_seletor_vazio_nao_quebra(self):
        processo = self.processo_em_analise()
        self.client.force_login(self.analista_lic)
        resposta = self.client.post(
            reverse('tram_direcionar_assinatura', args=[processo.id]),
            {'destinatario': ''})
        self.assertLess(resposta.status_code, 500)
        processo.refresh_from_db()
        self.assertEqual(processo.situacao_tramite, 'EM_ANALISE')

    def test_admin_exclui_processo_na_lista_de_ativos(self):
        processo = self.novo_processo(numero_processo='3003/2026')
        admin = self.gestao
        admin.is_superuser = True
        admin.save(update_fields=['is_superuser'])

        self.client.force_login(self.protocolo)
        negado = self.client.post(reverse('apagar_processo', args=[processo.id]))
        self.assertEqual(negado.status_code, 403)
        self.assertTrue(Processo.objects.filter(id=processo.id).exists())

        self.client.force_login(admin)
        pagina = self.client.get(reverse('listar_processos'))
        self.assertContains(pagina, 'Excluir processo')
        resposta = self.client.post(reverse('apagar_processo', args=[processo.id]))
        self.assertEqual(resposta.status_code, 200)
        self.assertFalse(Processo.objects.filter(id=processo.id).exists())

    def test_admin_apaga_usuario(self):
        from django.contrib.auth.models import User
        admin = self.gestao
        admin.is_staff = True
        admin.is_superuser = True
        admin.save()
        self.client.force_login(self.protocolo)
        self.client.post(reverse('apagar_usuario', args=[self.analista_lic2.id]))
        self.assertTrue(User.objects.filter(id=self.analista_lic2.id).exists())

        self.client.force_login(admin)
        pagina = self.client.get(reverse('manage_users'))
        self.assertContains(pagina, 'Apagar usuário')
        resposta = self.client.post(reverse('apagar_usuario', args=[self.analista_lic2.id]))
        self.assertRedirects(resposta, reverse('manage_users'))
        self.assertFalse(User.objects.filter(id=self.analista_lic2.id).exists())
        self.client.post(reverse('apagar_usuario', args=[admin.id]))
        self.assertTrue(User.objects.filter(id=admin.id).exists())

    def test_desativar_usuario_nao_apaga(self):
        admin = self.gestao
        admin.is_staff = True
        admin.is_superuser = True
        admin.save()
        self.client.force_login(admin)
        self.client.post(reverse('delete_user', args=[self.analista_lic2.id]))
        self.analista_lic2.refresh_from_db()
        self.assertFalse(self.analista_lic2.is_active)

    def test_admin_redefine_senha_de_outro_usuario(self):
        admin = self.gestao
        admin.is_staff = True
        admin.is_superuser = True
        admin.save()
        self.client.force_login(admin)
        nova = 'SenhaTemp#2026a'
        resposta = self.client.post(
            reverse('reset_user_password', args=[self.analista_lic2.id]),
            {'new_password1': nova, 'new_password2': nova})
        self.assertRedirects(resposta, reverse('manage_users'))
        self.analista_lic2.refresh_from_db()
        self.assertTrue(self.analista_lic2.check_password(nova))
        self.assertFalse(self.analista_lic2.check_password('senha-teste-123'))

    def test_gestao_devolve_processo_a_fila(self):
        processo = self.processo_em_analise()
        self.client.force_login(self.gestao)
        resposta = self.client.post(
            reverse('tram_declinar_analise', args=[processo.id]),
            {'motivo': 'Tentativa de teste'})
        self.assertRedirects(resposta, reverse('gestao_processos'))
        processo.refresh_from_db()
        self.assertEqual(processo.situacao_tramite, 'DISPONIVEL')
        self.assertIsNone(processo.analista_responsavel)

    def test_admin_desfaz_analise_e_encaminhamento_pela_tela(self):
        from .base import criar_usuario
        admin = criar_usuario('admin_undo', 'PROTOCOLO', is_superuser=True)
        processo = self.processo_em_analise()
        self.client.force_login(admin)
        fila = self.client.get(reverse('gestao_processos'))
        self.assertEqual(fila.status_code, 200)
        self.assertContains(fila, 'Devolver à fila')
        self.client.post(
            reverse('tram_desfazer_tramite', args=[processo.id]),
            {'motivo': 'Desfeito pelo administrador'})
        processo.refresh_from_db()
        self.assertEqual(processo.situacao_tramite, 'DISPONIVEL')

        processo = self.processo_em_analise()
        tramitacao.liberar_assinatura(processo.id, self.analista_lic)
        assinar = self.client.get(reverse('gestao_liberados_assinatura'))
        self.assertContains(assinar, 'Desfazer encaminhamento')
        self.client.post(
            reverse('tram_desfazer_tramite', args=[processo.id]),
            {'motivo': 'Desfeito pelo administrador'})
        processo.refresh_from_db()
        self.assertEqual(processo.situacao_tramite, 'EM_ANALISE')

    def test_nao_admin_nao_redefine_senha(self):
        self.client.force_login(self.protocolo)
        resposta = self.client.get(
            reverse('reset_user_password', args=[self.analista_lic2.id]))
        self.assertEqual(resposta.status_code, 302)

    def test_admin_nao_redefine_a_propria_senha_por_essa_tela(self):
        admin = self.gestao
        admin.is_staff = True
        admin.is_superuser = True
        admin.save()
        self.client.force_login(admin)
        resposta = self.client.get(reverse('reset_user_password', args=[admin.id]))
        self.assertRedirects(resposta, reverse('password_change'))

    def test_fluxo_completo_por_http(self):
        processo = self.processo_em_analise()
        self.client.force_login(self.analista_lic)
        self.client.post(reverse('tram_liberar_assinatura', args=[processo.id]))
        self.client.force_login(self.protocolo)
        self.client.post(reverse('tram_disponibilizar_retirada', args=[processo.id]))
        self.client.post(reverse('tram_registrar_saida'), {'processos': [processo.id]})
        processo.refresh_from_db()
        self.assertEqual(processo.situacao_tramite, 'SAIDA_CONCLUIDA',
                         'verificar nome do campo do lote em views_tramitacao.registrar_saida')
        self.assertFalse(tramitacao.ativos().filter(id=processo.id).exists())


@override_settings(**STATIC)
class ProcessosAtivosOrdemTest(BaseProcessoTestCase):

    def test_ordem_por_prioridade_e_entrada(self):
        normal_tarde = self.novo_processo(
            numero_processo='6101/2026',
            data_entrada='2026-09-10', hora_entrada='14:00')
        prioritario = self.novo_processo(
            numero_processo='6102/2026',
            data_entrada='2026-09-12', hora_entrada='09:00')
        tramitacao.alterar_prioridade(prioritario.id, self.gestao, 'PRIORITARIO')
        normal_cedo = self.novo_processo(
            numero_processo='6103/2026',
            data_entrada='2026-09-08', hora_entrada='08:00')
        urgente = self.novo_processo(
            numero_processo='6104/2026',
            data_entrada='2026-09-15', hora_entrada='18:00')
        tramitacao.alterar_prioridade(urgente.id, self.gestao, 'URGENTE')

        self.client.force_login(self.protocolo)
        resposta = self.client.get(reverse('listar_processos'))
        numeros = [p.numero_processo for p in resposta.context['processos']]
        # Ignora processos do setUp da base que possam aparecer na lista.
        ordem = [n for n in numeros if n.startswith('610')]
        self.assertEqual(ordem, [
            urgente.numero_processo,
            prioritario.numero_processo,
            normal_cedo.numero_processo,
            normal_tarde.numero_processo,
        ])


@override_settings(**STATIC)
class FilaAnalistaESituacaoTest(BaseProcessoTestCase):

    def test_processos_ativos_mostra_quem_analisa(self):
        processo = self.processo_em_analise()
        self.client.force_login(self.protocolo)
        resposta = self.client.get(reverse('listar_processos') + '?fase=pendentes')
        self.assertContains(resposta, 'Em análise por')
        self.assertContains(resposta, processo.nome_analista)

    def test_processos_ativos_separa_pendentes_e_finalizados(self):
        from processos_app.models import LinhaControleRelatorio

        pendente = self.novo_processo(numero_processo='7101/2026')
        finalizado = self.novo_processo(
            self.especie_liq, numero_processo='7102/2026')
        LinhaControleRelatorio.objects.create(
            processo=finalizado,
            numero_relatorio='7102',
            numero_processo=finalizado.numero_processo,
            data_relatorio=timezone.localdate(),
            situacao_linha='ATIVA',
            sequencia='LIQUIDACOES',
            grupo='LIQUIDACOES',
        )
        finalizado.numero_relatorio = '7102'
        finalizado.save(update_fields=['numero_relatorio'])

        self.client.force_login(self.protocolo)
        # Padrão Pendentes: o que ainda aguarda análise (não é saída física).
        padrao = self.client.get(reverse('listar_processos'))
        self.assertEqual(padrao.context['fase'], 'pendentes')
        self.assertContains(padrao, pendente.numero_processo)
        self.assertNotContains(padrao, finalizado.numero_processo)

        pend = self.client.get(reverse('listar_processos') + '?fase=pendentes')
        self.assertContains(pend, 'Pendentes')
        self.assertContains(pend, pendente.numero_processo)
        self.assertNotContains(pend, finalizado.numero_processo)

        fim = self.client.get(reverse('listar_processos') + '?fase=finalizados')
        self.assertContains(fim, finalizado.numero_processo)
        self.assertNotContains(fim, pendente.numero_processo)

    def test_reentrada_nao_herda_finalizados_da_passagem_anterior(self):
        """Análise antiga do mesmo número não manda a reentrada para Finalizados."""
        from processos_app.models import LinhaControleRelatorio
        from processos_app.services import relatorios as svc_relatorios

        antiga = self.novo_processo(
            self.especie_liq, numero_processo='15371/2021')
        LinhaControleRelatorio.objects.create(
            processo=antiga,
            numero_relatorio='88001',
            numero_processo=antiga.numero_processo,
            data_relatorio=timezone.localdate(),
            situacao_linha='ATIVA',
            sequencia='LIQUIDACOES',
            grupo='LIQUIDACOES',
        )
        antiga.numero_relatorio = '88001'
        antiga.save(update_fields=['numero_relatorio'])
        tramitacao.registrar_saida_direta(antiga.id, self.protocolo)

        reentrada = self.novo_processo(
            self.especie_liq, numero_processo='15371/2021', volume='4')

        self.client.force_login(self.protocolo)
        pend = self.client.get(reverse('listar_processos') + '?fase=pendentes')
        self.assertContains(pend, '15371/2021')
        fim = self.client.get(reverse('listar_processos') + '?fase=finalizados')
        self.assertNotContains(fim, '15371/2021')
        self.assertFalse(
            svc_relatorios.processo_tem_analise_registrada(reentrada))
        self.assertIn(
            reentrada.id,
            tramitacao.pendentes_como_ativos().values_list('id', flat=True))

    def test_saida_remove_de_ativos_e_da_fila_analista(self):
        processo = self.novo_processo(
            self.especie_liq, numero_processo='saida-fila/2026')
        self.client.force_login(self.protocolo)
        ativos = self.client.get(reverse('listar_processos'))
        self.assertContains(ativos, 'saida-fila/2026')

        self.client.force_login(self.analista_liq)
        fila = self.client.get(reverse('area_analista'))
        self.assertContains(fila, 'saida-fila/2026')

        tramitacao.registrar_saida_direta(processo.id, self.protocolo)

        self.client.force_login(self.protocolo)
        ativos2 = self.client.get(reverse('listar_processos'))
        self.assertNotContains(ativos2, 'saida-fila/2026')

        self.client.force_login(self.analista_liq)
        fila2 = self.client.get(reverse('area_analista'))
        self.assertNotContains(fila2, 'saida-fila/2026')

    def test_fila_nao_mostra_disponivel_ja_analisado(self):
        """DISPONÍVEL com análise no Controle some da fila (como de Pendentes)."""
        from processos_app.models import LinhaControleRelatorio

        fantasma = self.novo_processo(
            self.especie_liq, numero_processo='fantasma/2026')
        LinhaControleRelatorio.objects.create(
            processo=fantasma,
            numero_relatorio='99001',
            numero_processo=fantasma.numero_processo,
            data_relatorio=timezone.localdate(),
            situacao_linha='ATIVA',
            sequencia='LIQUIDACOES',
            grupo='LIQUIDACOES',
        )
        fantasma.numero_relatorio = '99001'
        fantasma.save(update_fields=['numero_relatorio'])

        self.client.force_login(self.protocolo)
        pend = self.client.get(reverse('listar_processos') + '?fase=pendentes')
        self.assertNotContains(pend, 'fantasma/2026')

        self.client.force_login(self.analista_liq)
        fila = self.client.get(reverse('area_analista'))
        self.assertNotContains(fila, 'fantasma/2026')

        self.client.force_login(self.gestao)
        gestao = self.client.get(reverse('gestao_processos'))
        self.assertNotContains(gestao, 'fantasma/2026')

    def test_processo_com_saida_inconsistente_nao_aparece_na_fila(self):
        """Saída gravada em qualquer campo tira de Ativos e da fila."""
        inconsistente = self.novo_processo(
            self.especie_liq, numero_processo='saiu-sujo/2026')
        # Situação ainda "ativa", mas já tem marca de saída.
        inconsistente.saida_concluida_em = timezone.now()
        inconsistente.data_saida = timezone.localdate()
        inconsistente.save(update_fields=['saida_concluida_em', 'data_saida'])

        self.assertFalse(
            tramitacao.ativos().filter(id=inconsistente.id).exists())

        self.client.force_login(self.protocolo)
        ativos = self.client.get(reverse('listar_processos') + '?fase=todos')
        self.assertNotContains(ativos, 'saiu-sujo/2026')

        self.client.force_login(self.analista_liq)
        fila = self.client.get(reverse('area_analista'))
        self.assertNotContains(fila, 'saiu-sujo/2026')

    def test_fila_padrao_mostra_quem_esta_analisando_e_esconde_o_controlador(self):
        disponivel = self.novo_processo(numero_processo='1001/2026')
        comigo = self.processo_em_analise()
        comigo.numero_processo = '1002/2026'
        comigo.save(update_fields=['numero_processo'])
        outro = self.processo_em_analise(analista=self.analista_lic2)
        outro.numero_processo = '1003/2026'
        outro.save(update_fields=['numero_processo'])
        no_controlador = self.processo_em_analise()
        no_controlador.numero_processo = '1004/2026'
        no_controlador.save(update_fields=['numero_processo'])
        tramitacao.liberar_assinatura(no_controlador.id, self.analista_lic)

        self.client.force_login(self.analista_lic)
        resposta = self.client.get(reverse('area_analista'))
        # Com o Controlador = Finalizados em Ativos — fora da fila (Pendentes).
        self.assertEqual(resposta.context['total_com_controlador'], 0)
        self.assertNotContains(resposta, '1004/2026')
        self.assertContains(resposta, '1001/2026')
        self.assertContains(resposta, '1002/2026')
        self.assertContains(resposta, '1003/2026')
        self.assertContains(resposta, outro.nome_analista)
        html_acoes = resposta.content.decode()
        trecho_outro = html_acoes.split('1003/2026', 1)[1].split('</tr>', 1)[0]
        self.assertIn(outro.nome_analista, trecho_outro)
        self.assertNotIn('Assumir processo', trecho_outro)
        self.assertNotContains(resposta, '1004/2026')

        # Liberados também fora de Pendentes — some da fila.
        resposta = self.client.get(reverse('area_analista') + '?filtro=liberados')
        self.assertNotContains(resposta, '1004/2026')

    def test_fila_analista_somente_consulta_sem_assumir_nem_analisar(self):
        disponivel = self.novo_processo(numero_processo='fila-disp/2026')
        comigo = self.processo_em_analise()
        comigo.numero_processo = 'fila-meu/2026'
        comigo.save(update_fields=['numero_processo'])

        self.client.force_login(self.analista_lic)
        pagina = self.client.get(reverse('area_analista'))
        self.assertTrue(pagina.context['fila_somente_leitura'])
        self.assertFalse(pagina.context['pode_assumir'])
        self.assertContains(pagina, 'Pendentes de Processos Ativos')
        self.assertContains(pagina, 'fila-disp/2026')
        self.assertContains(pagina, 'fila-meu/2026')
        self.assertNotContains(pagina, 'Assumir processo')
        self.assertNotContains(pagina, 'Encaminhar ao Controlador')
        self.assertNotContains(pagina, 'Devolver à fila')
        # Só "Ver" nas ações — não o botão "Analisar" da fila.
        self.assertNotRegex(
            pagina.content.decode(),
            r'edit_note[\s\S]{0,80}Analisar')
        self.assertContains(pagina, 'Ver')

        # Assumir continua pela janela "Assumir" (Nova análise), não pela fila.
        assumiu = self.client.post(
            reverse('assumir_processo', args=[disponivel.id]))
        self.assertRedirects(
            assumiu, reverse('analista_processo', args=[disponivel.id]))
        disponivel.refresh_from_db()
        self.assertEqual(disponivel.situacao_tramite, 'EM_ANALISE')
        self.assertEqual(disponivel.analista_responsavel_id, self.analista_lic.id)

    def test_fila_analista_so_mostra_entrada_do_protocolo(self):
        """Fila de análise: só processos cadastrados pelo Protocolo."""
        from processos_app.models import EventoProcesso, Processo
        from processos_app.services import relatorios as svc_relatorios

        do_protocolo = self.novo_processo(
            self.especie_liq, numero_processo='fila-proto/2026')
        sem_protocolo = svc_relatorios.salvar_nova_analise(self.analista_liq, {
            'numero_processo': 'fila-nova/2026',
            'grupo': 'LIQUIDACOES',
            'gerar_numero': '1',
            'data_relatorio': '2026-10-08',
            'objeto': 'Stub Nova análise',
            'analista': 'Caio Liq',
        })
        stub_id = sem_protocolo['processo_id']
        self.assertFalse(
            EventoProcesso.objects.filter(
                processo_id=stub_id, tipo='PROCESSO_CADASTRADO').exists())
        self.assertTrue(Processo.objects.filter(id=stub_id).exists())
        self.assertTrue(Processo.objects.filter(id=do_protocolo.id).exists())

        self.client.force_login(self.analista_liq)
        pagina = self.client.get(reverse('area_analista'))
        self.assertEqual(pagina.status_code, 200)
        self.assertContains(pagina, 'fila-proto/2026')
        self.assertNotContains(pagina, 'fila-nova/2026')

    def test_historico_mostra_usuario_entrada_e_saida(self):
        processo = self.novo_processo(numero_processo='hist-user/2026')
        self.anexar_teste(processo)
        tramitacao.assumir(processo.id, self.analista_lic)
        tramitacao.liberar_assinatura(processo.id, self.analista_lic)
        processo.refresh_from_db()
        processo.situacao_tramite = 'DISPONIVEL_RETIRADA'
        processo.save(update_fields=['situacao_tramite'])
        tramitacao.registrar_saida([processo.id], self.protocolo)

        self.client.force_login(self.protocolo)
        historico = self.client.get(
            reverse('ver_historico_processo', args=[processo.id]))
        self.assertContains(historico, 'Entrada por')
        self.assertContains(historico, 'Saída por')
        self.assertContains(historico, self.protocolo.get_full_name()
                            or self.protocolo.username)
        passagem = historico.context['passagens'][0]
        self.assertEqual(passagem['entrada_por'],
                         self.protocolo.get_full_name() or self.protocolo.username)
        self.assertEqual(passagem['saida_por'],
                         self.protocolo.get_full_name() or self.protocolo.username)

    def test_fila_liquidacoes_separa_por_blocos_de_sequencia(self):
        especie_bolsa, _ = EspecieProcesso.objects.get_or_create(
            nome='Concessão Aux. Bolsa Atleta', grupo=LIQ,
            defaults={'ativo': True, 'gera_relatorio': True,
                      'sequencia_numeracao': 'BOLSA_ATLETA'})
        if especie_bolsa.sequencia_numeracao != 'BOLSA_ATLETA':
            especie_bolsa.sequencia_numeracao = 'BOLSA_ATLETA'
            especie_bolsa.gera_relatorio = True
            especie_bolsa.ativo = True
            especie_bolsa.save(update_fields=[
                'sequencia_numeracao', 'gera_relatorio', 'ativo'])
        liquidacao = self.novo_processo(
            especie=self.especie_liq, numero_processo='2001/2026')
        bolsa = self.novo_processo(
            especie=especie_bolsa, numero_processo='2002/2026')

        self.client.force_login(self.analista_liq)
        pagina = self.client.get(reverse('area_analista'))
        self.assertContains(pagina, 'Liquidação')
        self.assertContains(pagina, 'Bolsa Atleta')
        self.assertContains(pagina, '2001/2026')
        self.assertContains(pagina, '2002/2026')
        grupos = pagina.context['grupos_liquidacoes']
        self.assertEqual([g['codigo'] for g in grupos], ['LIQUIDACOES', 'BOLSA_ATLETA'])
        self.assertEqual(grupos[0]['processos'][0].id, liquidacao.id)
        self.assertEqual(grupos[1]['processos'][0].id, bolsa.id)

    def test_analise_bolsa_atleta_mostra_processo_de_prestacao(self):
        especie_pc, _ = EspecieProcesso.objects.get_or_create(
            nome='P.C. Bolsa Atleta', grupo=LIQ,
            defaults={'ativo': True, 'gera_relatorio': True,
                      'sequencia_numeracao': 'BOLSA_ATLETA'})
        if especie_pc.sequencia_numeracao != 'BOLSA_ATLETA':
            especie_pc.sequencia_numeracao = 'BOLSA_ATLETA'
            especie_pc.gera_relatorio = True
            especie_pc.ativo = True
            especie_pc.save(update_fields=[
                'sequencia_numeracao', 'gera_relatorio', 'ativo'])
        processo = self.novo_processo(
            especie=especie_pc, numero_processo='24099/2025')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.refresh_from_db()

        self.client.force_login(self.analista_liq)
        tela = self.client.get(reverse('analista_processo', args=[processo.id]))
        self.assertEqual(tela.status_code, 200)
        self.assertContains(tela, 'Processo de prestação')
        self.assertContains(tela, 'Processo de concessão')
        self.assertContains(tela, 'id="processo_prestacao"')
        self.assertContains(tela, '24099/2025')

        rel = tela.context['processos_relacionados']
        # Concessão = nº do processo; prestação vazia; os dois editáveis.
        self.assertEqual(rel['concessao'], '24099/2025')
        self.assertEqual(rel['prestacao'], '')
        self.assertEqual(rel['concessao_name'], 'numero_processo')
        self.assertEqual(rel['prestacao_name'], 'processo_prestacao')
        self.assertContains(tela, 'name="numero_processo"')
        self.assertContains(tela, 'name="processo_prestacao"')

        resp = self.client.post(reverse('analista_processo', args=[processo.id]), {
            'valor': 'R$ 3.000,00',
            'destino': 'Unidade de Teste',
            'periodo': 'Aline Resp',
            'contratada': 'Laura Atleta',
            'objeto': 'Prestação 2º semestre — Jiu-Jitsu',
            'numero_processo': '24099/2025',
            'processo_prestacao': '24.099/2025',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            'observacao': '',
        })
        self.assertEqual(resp.status_code, 302)
        processo.refresh_from_db()
        self.assertEqual(processo.numero_processo, '24099/2025')
        self.assertEqual(processo.processo_prestacao, '24.099/2025')
        from processos_app.services import relatorios
        dados_linha = relatorios._dados_da_linha(processo)
        self.assertEqual(dados_linha['volume'], '24099/2025')
        self.assertEqual(dados_linha['numero_processo'], '24.099/2025')

        especie_conc, _ = EspecieProcesso.objects.get_or_create(
            nome='Concessão Aux. Bolsa Atleta', grupo=LIQ,
            defaults={'ativo': True, 'gera_relatorio': True,
                      'sequencia_numeracao': 'BOLSA_ATLETA'})
        if especie_conc.sequencia_numeracao != 'BOLSA_ATLETA':
            especie_conc.sequencia_numeracao = 'BOLSA_ATLETA'
            especie_conc.save(update_fields=['sequencia_numeracao'])
        conc = self.novo_processo(
            especie=especie_conc, numero_processo='22140/2024')
        tramitacao.assumir(conc.id, self.analista_liq)
        tela_conc = self.client.get(reverse('analista_processo', args=[conc.id]))
        rel_conc = tela_conc.context['processos_relacionados']
        self.assertEqual(rel_conc['concessao'], '22140/2024')
        self.assertEqual(rel_conc['prestacao'], '')
        resp_conc = self.client.post(reverse('analista_processo', args=[conc.id]), {
            'valor': 'R$ 3.000,00',
            'destino': 'Unidade de Teste',
            'periodo': 'Responsavel',
            'contratada': 'Atleta',
            'objeto': 'Concessão — Natação',
            'numero_processo': '22140/2024',
            'processo_prestacao': '',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            'observacao': '',
        })
        self.assertEqual(resp_conc.status_code, 302)
        conc.refresh_from_db()
        self.assertEqual(conc.processo_prestacao, '')
        self.assertEqual(relatorios._dados_da_linha(conc)['volume'], '22140/2024')
        self.assertEqual(relatorios._dados_da_linha(conc)['numero_processo'], '')
