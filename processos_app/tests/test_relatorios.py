from django.core.exceptions import PermissionDenied
from django.urls import reverse
from django.utils import timezone

from processos_app.models import LinhaControleRelatorio
from processos_app.services import processos as svc_processos
from processos_app.services import relatorios, tramitacao

from .base import BaseProcessoTestCase, criar_usuario

import datetime


class ControleRelatorioTest(BaseProcessoTestCase):

    def _salvar_liquidacao(self, processo, analista=None, gerar_numero=True):
        analista = analista or self.analista_liq
        if processo.situacao_tramite != 'EM_ANALISE':
            tramitacao.assumir(processo.id, analista)
            processo.refresh_from_db()
        dados = {
            'destino': 'Unidade de Teste',
            'valor': '1000',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
        }
        if gerar_numero:
            dados['gerar_numero'] = '1'
        svc_processos.aplicar_analise(processo, dados, analista)
        processo.refresh_from_db()
        return processo

    def test_busca_encontra_por_qualquer_campo(self):
        processo = self.novo_processo(self.especie_liq, numero_processo='busca-1/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.refresh_from_db()
        svc_processos.aplicar_analise(processo, {
            'destino': 'Secretaria da Cultura',
            'valor': '2500,50',
            'periodo': 'Jan/2026',
            'observacao': 'Obs exclusiva busca',
            'status_analise': 'PROSSEGUIMENTO_COM_RESSALVA',
            'gerar_numero': '1',
        }, self.analista_liq)
        processo.refresh_from_db()
        linha = LinhaControleRelatorio.objects.get(processo=processo)

        self.client.force_login(self.analista_liq)
        for termo in (
            processo.numero_relatorio,
            'busca-1/2026',
            'Secretaria da Cultura',
            'Jan/2026',
            'Obs exclusiva busca',
            'com ressalva',
            linha.data_relatorio.strftime('%d/%m/%Y'),
        ):
            pagina = self.client.get(
                reverse('controle_relatorio') + f'?termo={termo}')
            self.assertContains(
                pagina, 'busca-1/2026',
                msg_prefix=f'termo={termo!r}')

        vazia = self.client.get(
            reverse('controle_relatorio') + '?termo=nao-existe-xyz')
        self.assertNotContains(vazia, 'busca-1/2026')

    def test_listar_aceita_numero_relatorio_nao_inteiro(self):
        """Planilha antiga pode ter nº tipo '56 A'; Postgres não pode Cast cego."""
        admin = criar_usuario('admin_num_txt', 'GESTAO', is_superuser=True)
        LinhaControleRelatorio.objects.create(
            processo=None,
            numero_relatorio='56 A',
            numero_processo='x/2026',
            data_relatorio=datetime.date(2026, 1, 10),
            situacao_linha='HISTORICA',
            grupo='LIQUIDACOES',
            sequencia='LIQUIDACOES',
        )
        LinhaControleRelatorio.objects.create(
            processo=None,
            numero_relatorio='100',
            numero_processo='y/2026',
            data_relatorio=datetime.date(2026, 1, 11),
            situacao_linha='HISTORICA',
            grupo='LIQUIDACOES',
            sequencia='LIQUIDACOES',
        )
        lista = list(relatorios.listar(admin, 'LIQUIDACOES'))
        numeros = [linha.numero_relatorio for linha in lista]
        self.assertIn('56 A', numeros)
        self.assertIn('100', numeros)
        self.client.force_login(admin)
        pagina = self.client.get(
            reverse('controle_relatorio') + '?aba=LIQUIDACOES&secao=analises')
        self.assertEqual(pagina.status_code, 200)
        self.assertContains(pagina, '56 A')

    def test_desvincular_processo_mantem_numero_na_planilha(self):
        admin = criar_usuario('admin_desvinc', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 7000)
        processo = self.novo_processo(
            self.especie_liq, numero_processo='desv-1/2026')
        self._salvar_liquidacao(processo)
        self.assertEqual(processo.numero_relatorio, '7001')
        linha = LinhaControleRelatorio.objects.get(processo=processo)

        self.client.force_login(admin)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, 'Desvincular processo deste número')

        resp = self.client.post(
            reverse('controle_relatorio_desvincular_linha', args=[linha.id]), {
                'grupo': 'LIQUIDACOES',
            })
        self.assertEqual(resp.status_code, 302)
        linha.refresh_from_db()
        processo.refresh_from_db()
        self.assertIsNone(linha.processo_id)
        self.assertEqual(linha.numero_relatorio, '7001')
        self.assertEqual(linha.numero_processo, 'desv-1/2026')
        self.assertEqual(linha.situacao_linha, 'HISTORICA')
        self.assertFalse(processo.numero_relatorio)
        self.assertIn('desvinculado', (linha.observacao or '').casefold())

        planilha = self.client.get(reverse('controle_relatorio'))
        self.assertContains(planilha, 'data-vincular-linha')

    def test_vincular_processo_finalizado_preenche_historico(self):
        admin = criar_usuario('admin_vinc_fim', 'GESTAO', is_superuser=True)
        processo = self.novo_processo(
            self.especie_liq, numero_processo='fim-1/2026',
            secretaria='Secretaria Finalizada',
            objeto='Objeto da última passagem',
            contratada='Empresa X',
            valor='R$ 10,00',
        )
        processo.situacao_tramite = 'SAIDA_CONCLUIDA'
        processo.data_saida = datetime.date(2026, 2, 1)
        processo.tecnico = 'Analista Z'
        processo.save()

        linha = LinhaControleRelatorio.objects.create(
            processo=None,
            numero_relatorio='888',
            numero_processo='',
            data_relatorio=datetime.date(2026, 1, 10),
            situacao_linha='HISTORICA',
            grupo='LIQUIDACOES',
            sequencia='LIQUIDACOES',
        )
        self.client.force_login(admin)
        resp = self.client.post(
            reverse('controle_relatorio_vincular_linha', args=[linha.id]), {
                'grupo': 'LIQUIDACOES',
                'numero_processo': 'fim-1/2026',
            })
        self.assertEqual(resp.status_code, 302)
        linha.refresh_from_db()
        processo.refresh_from_db()
        self.assertEqual(linha.numero_processo, 'fim-1/2026')
        self.assertEqual(linha.secretaria, 'Secretaria Finalizada')
        self.assertEqual(linha.objeto, 'Objeto da última passagem')
        self.assertEqual(linha.contratada, 'Empresa X')
        self.assertIsNone(linha.processo_id)
        self.assertEqual(linha.situacao_linha, 'HISTORICA')
        # Processo finalizado não recebe o número de relatório.
        self.assertFalse(processo.numero_relatorio)

    def test_editar_linha_historica_pela_planilha(self):
        admin = criar_usuario('admin_edit_hist', 'GESTAO', is_superuser=True)
        linha = LinhaControleRelatorio.objects.create(
            processo=None,
            numero_relatorio='777',
            numero_processo='old/2026',
            data_relatorio=datetime.date(2026, 1, 1),
            secretaria='Sec antiga',
            objeto='Objeto antigo',
            situacao_linha='HISTORICA',
            grupo='LIQUIDACOES',
            sequencia='LIQUIDACOES',
            observacao='Importado',
        )
        self.client.force_login(admin)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, 'Editar linha da planilha')
        resp = self.client.post(
            reverse('controle_relatorio_editar_linha', args=[linha.id]), {
                'grupo': 'LIQUIDACOES',
                'numero_relatorio': '777',
                'data_relatorio': '2026-03-15',
                'numero_processo': 'novo/2026',
                'secretaria': 'Sec nova',
                'objeto': 'Objeto novo',
                'valor': '1500',
                'analista': 'Ana',
                'sem_relatorio': '1',
            })
        self.assertEqual(resp.status_code, 302)
        linha.refresh_from_db()
        self.assertEqual(str(linha.data_relatorio), '2026-03-15')
        self.assertEqual(linha.numero_processo, 'novo/2026')
        self.assertEqual(linha.secretaria, 'Sec nova')
        self.assertEqual(linha.objeto, 'Objeto novo')
        self.assertTrue(linha.sem_relatorio)
        self.assertTrue(linha.valor.startswith('R$'))

    def test_analista_troca_numero_relatorio_pela_planilha(self):
        processo = self.novo_processo(self.especie_liq, numero_processo='troca-n/2026')
        self._salvar_liquidacao(processo)
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        numero_antigo = linha.numero_relatorio
        novo = str(int(numero_antigo) + 50)

        self.client.force_login(self.analista_liq)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, 'name="numero_relatorio"')
        resp = self.client.post(
            reverse('controle_relatorio_editar_linha', args=[linha.id]), {
                'grupo': 'LIQUIDACOES',
                'numero_relatorio': novo,
                'data_relatorio': '2026-10-06',
                'numero_processo': processo.numero_processo,
                'secretaria': processo.secretaria,
                'objeto': processo.objeto,
                'valor': '100',
                'analista': 'Analista',
            })
        self.assertEqual(resp.status_code, 302)
        linha.refresh_from_db()
        processo.refresh_from_db()
        self.assertEqual(linha.numero_relatorio, novo)
        self.assertEqual(processo.numero_relatorio, novo)

    def test_marcar_sem_relatorio_pela_planilha_do_controle(self):
        admin = criar_usuario('admin_sr_planilha', 'GESTAO', is_superuser=True)
        processo = self.novo_processo(self.especie_liq, numero_processo='sr-plan/2026')
        self._salvar_liquidacao(processo)
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertFalse(linha.sem_relatorio)

        self.client.force_login(admin)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, 'Marcar como sem relatório')
        resp = self.client.post(
            reverse('controle_relatorio_sem_relatorio', args=[linha.id]),
            {'grupo': 'LIQUIDACOES'},
        )
        self.assertEqual(resp.status_code, 302)
        linha.refresh_from_db()
        processo.refresh_from_db()
        self.assertTrue(linha.sem_relatorio)
        self.assertTrue(processo.sem_relatorio)

        resp = self.client.post(
            reverse('controle_relatorio_sem_relatorio', args=[linha.id]),
            {'grupo': 'LIQUIDACOES'},
        )
        self.assertEqual(resp.status_code, 302)
        linha.refresh_from_db()
        processo.refresh_from_db()
        self.assertFalse(linha.sem_relatorio)
        self.assertFalse(processo.sem_relatorio)

    def test_salvar_analise_sem_relatorio_aparece_no_controle(self):
        processo = self.novo_processo(self.especie_liq, numero_processo='sr-1/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.refresh_from_db()
        svc_processos.aplicar_analise(processo, {
            'destino': 'Unidade de Teste',
            'valor': '1000',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            'sem_relatorio': '1',
            'gerar_numero': '1',
        }, self.analista_liq)
        processo.refresh_from_db()
        self.assertTrue(processo.sem_relatorio)
        self.assertTrue(processo.numero_relatorio)
        self.assertTrue(processo.data_analise)
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertTrue(linha.sem_relatorio)
        self.assertEqual(linha.numero_relatorio, processo.numero_relatorio)
        self.assertEqual(linha.numero_processo, 'sr-1/2026')

        self.client.force_login(self.analista_liq)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, 'sr-1/2026')
        self.assertContains(pagina, 'tr-sem-relatorio')

        tela = self.client.get(reverse('analista_processo', args=[processo.id]))
        self.assertContains(tela, 'name="sem_relatorio"')
        self.assertContains(tela, 'Sem relatório')

    def test_admin_marca_sem_relatorio_em_processo_ja_numerado(self):
        admin = criar_usuario('admin_sr', 'GESTAO', is_superuser=True)
        processo = self.novo_processo(self.especie_liq, numero_processo='sr-adm/2026')
        self._salvar_liquidacao(processo)
        self.assertTrue(processo.numero_relatorio)

        self.client.force_login(admin)
        tela = self.client.get(
            reverse('analista_processo', args=[processo.id])
            + f'?next={reverse("controle_relatorio")}')
        self.assertTrue(tela.context['pode_editar'])
        self.assertContains(tela, 'name="sem_relatorio"')
        self.assertContains(tela, 'Edição do administrador')
        numero_antes = processo.numero_relatorio

        resposta = self.client.post(
            reverse('analista_processo', args=[processo.id]), {
                'destino': 'Unidade de Teste',
                'valor': '1000',
                'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
                'sem_relatorio': '1',
                'next': reverse('controle_relatorio'),
            })
        self.assertRedirects(resposta, reverse('controle_relatorio'))
        processo.refresh_from_db()
        self.assertTrue(processo.sem_relatorio)
        self.assertEqual(processo.numero_relatorio, numero_antes)
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertTrue(linha.sem_relatorio)
        self.assertEqual(linha.numero_relatorio, numero_antes)
        self.assertEqual(
            LinhaControleRelatorio.objects.filter(
                numero_relatorio=numero_antes).count(), 1)

        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, 'sr-adm/2026')
        self.assertContains(pagina, numero_antes)
        self.assertContains(pagina, 'tr-sem-relatorio')

        self.client.force_login(self.gestao)
        negado = self.client.post(
            reverse('analista_processo', args=[processo.id]), {
                'sem_relatorio': '1',
                'destino': 'Unidade de Teste',
                'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            })
        self.assertEqual(negado.status_code, 403)

    def test_nao_libera_liquidacao_sem_numero_relatorio(self):
        from processos_app.services.tramitacao import TransicaoInvalida
        processo = self.novo_processo(self.especie_liq, numero_processo='sr-lib/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.refresh_from_db()
        svc_processos.aplicar_analise(processo, {
            'destino': 'Unidade de Teste',
            'valor': '1000',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
        }, self.analista_liq)
        processo.refresh_from_db()
        self.assertFalse(processo.numero_relatorio)
        with self.assertRaises(TransicaoInvalida) as ctx:
            tramitacao.liberar_assinatura(processo.id, self.analista_liq)
        self.assertIn('Número do relatório', str(ctx.exception))

    def test_admin_informa_numero_em_processo_sem_relatorio(self):
        admin = criar_usuario('admin_num', 'GESTAO', is_superuser=True)
        processo = self.novo_processo(self.especie_liq, numero_processo='sr-num/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.refresh_from_db()
        svc_processos.aplicar_analise(processo, {
            'destino': 'Unidade de Teste',
            'valor': '1000',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            'sem_relatorio': '1',
        }, self.analista_liq)
        processo.refresh_from_db()
        self.assertTrue(processo.sem_relatorio)
        self.assertFalse(processo.numero_relatorio)

        self.client.force_login(admin)
        resposta = self.client.post(
            reverse('analista_processo', args=[processo.id]), {
                'destino': 'Unidade de Teste',
                'valor': '1000',
                'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
                'numero_relatorio': '1912',
                'data_relatorio': '2026-10-01',
                'next': reverse('controle_relatorio'),
            })
        self.assertRedirects(resposta, reverse('controle_relatorio'))
        processo.refresh_from_db()
        self.assertEqual(processo.numero_relatorio, '1912')
        self.assertFalse(processo.sem_relatorio)
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertEqual(linha.numero_relatorio, '1912')
        self.assertFalse(linha.sem_relatorio)

        resposta = self.client.post(
            reverse('analista_processo', args=[processo.id]), {
                'destino': 'Unidade de Teste',
                'valor': '1000',
                'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
                'sem_relatorio': '1',
                'numero_relatorio': '1912',
                'data_relatorio': '2026-10-01',
            })
        self.assertEqual(resposta.status_code, 302)
        processo.refresh_from_db()
        self.assertEqual(processo.numero_relatorio, '1912')
        self.assertTrue(processo.sem_relatorio)

    def test_admin_define_ultimo_e_proximo_salvo_continua(self):
        admin = criar_usuario('admin_rel', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 2000)
        estado = relatorios.estado_sequencia()
        self.assertEqual(estado['ultimo'], 2000)
        self.assertEqual(estado['proximo'], 2001)

        processo = self.novo_processo(self.especie_liq, numero_processo='9/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.refresh_from_db()
        self.assertFalse(processo.numero_relatorio)
        self.assertFalse(LinhaControleRelatorio.objects.filter(processo=processo).exists())

        self._salvar_liquidacao(processo)
        self.assertEqual(processo.numero_relatorio, '2001')
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertEqual(linha.numero_processo, '9/2026')
        self.assertEqual(linha.numero_relatorio, '2001')
        self.assertEqual(linha.analista, processo.nome_analista)

    def test_analista_nao_define_ultimo(self):
        with self.assertRaises(PermissionDenied):
            relatorios.definir_ultimo_numero(self.analista_liq, 10)

    def test_salvar_analise_nao_gera_numero_sozinho(self):
        processo = self.novo_processo(self.especie_liq, numero_processo='10/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.refresh_from_db()
        self.assertFalse(processo.numero_relatorio)
        svc_processos.aplicar_analise(processo, {
            'destino': 'Unidade de Teste',
            'valor': '1000',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
        }, self.analista_liq)
        processo.refresh_from_db()
        self.assertFalse(processo.numero_relatorio)
        self.assertFalse(
            LinhaControleRelatorio.objects.filter(processo=processo).exists())
        self.assertEqual(processo.destino, 'Unidade de Teste')

    def test_gerar_numero_emite_sequencial_e_planilha(self):
        processo = self.novo_processo(self.especie_liq, numero_processo='10b/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.refresh_from_db()
        self.assertFalse(processo.numero_relatorio)
        svc_processos.aplicar_analise(processo, {
            'destino': 'Unidade de Teste',
            'valor': '1000',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            'gerar_numero': '1',
        }, self.analista_liq)
        processo.refresh_from_db()
        self.assertTrue(processo.numero_relatorio)
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertEqual(linha.destino, 'Unidade de Teste')
        self.assertEqual(linha.valor, 'R$ 1.000,00')
        self.assertEqual(linha.status_analise, 'PROSSEGUIMENTO_SEM_RESSALVA')
        numero = processo.numero_relatorio
        tramitacao.liberar_assinatura(processo.id, self.analista_liq)
        processo.refresh_from_db()
        self.assertEqual(processo.numero_relatorio, numero)

        self.client.force_login(self.analista_liq)
        pagina = self.client.get(reverse('analista_processo', args=[processo.id]))
        # Já tem número: botão Gerar some.
        self.assertNotContains(pagina, 'value="gerar_numero"')

        sem_num = self.novo_processo(self.especie_liq, numero_processo='10c/2026')
        tramitacao.assumir(sem_num.id, self.analista_liq)
        tela = self.client.get(reverse('analista_processo', args=[sem_num.id]))
        self.assertContains(tela, 'value="gerar_numero"')
        self.assertContains(tela, 'Gerar número')
        self.assertContains(tela, 'Salvar análise')

    def test_declinar_devolve_o_numero_ao_proximo_analista(self):
        admin = criar_usuario('admin_gasto', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 2000)
        primeiro = self.novo_processo(self.especie_liq, numero_processo='11/2026')
        self._salvar_liquidacao(primeiro)
        self.assertEqual(primeiro.numero_relatorio, '2001')
        tramitacao.declinar_analise(primeiro.id, self.analista_liq, 'Teste')
        primeiro.refresh_from_db()
        self.assertFalse(primeiro.numero_relatorio)
        liberado = LinhaControleRelatorio.objects.get(numero_relatorio='2001')
        self.assertIsNone(liberado.processo_id)
        self.assertEqual(liberado.situacao_linha, 'RESERVADA')
        self.assertEqual(liberado.data_relatorio, timezone.localdate())

        segundo = self.novo_processo(self.especie_liq, numero_processo='11b/2026')
        self._salvar_liquidacao(segundo)
        self.assertEqual(segundo.numero_relatorio, '2001')
        liberado_id = liberado.id
        self.assertFalse(
            LinhaControleRelatorio.objects.filter(id=liberado_id).exists())
        ativo = LinhaControleRelatorio.objects.get(processo=segundo)
        self.assertEqual(ativo.numero_relatorio, '2001')
        self.assertEqual(ativo.situacao_linha, 'ATIVA')

    def test_declinar_no_meio_reusa_numero_verde_no_mesmo_dia(self):
        admin = criar_usuario('admin_meio', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 2000)
        primeiro = self.novo_processo(self.especie_liq, numero_processo='11c/2026')
        segundo = self.novo_processo(self.especie_liq, numero_processo='11d/2026')
        self._salvar_liquidacao(primeiro)
        self._salvar_liquidacao(segundo)
        self.assertEqual(primeiro.numero_relatorio, '2001')
        self.assertEqual(segundo.numero_relatorio, '2002')
        tramitacao.declinar_analise(primeiro.id, self.analista_liq, 'Teste')

        liberado = LinhaControleRelatorio.objects.get(numero_relatorio='2001')
        self.assertEqual(liberado.situacao_linha, 'RESERVADA')
        self.assertIsNone(liberado.processo_id)

        # No mesmo dia, o 2001 verde é reaproveitado antes do contador (2003).
        terceiro = self.novo_processo(self.especie_liq, numero_processo='11e/2026')
        self._salvar_liquidacao(terceiro)
        self.assertEqual(terceiro.numero_relatorio, '2001')
        quarto = self.novo_processo(self.especie_liq, numero_processo='11f/2026')
        self._salvar_liquidacao(quarto)
        self.assertEqual(quarto.numero_relatorio, '2003')

    def test_declinar_numero_verde_de_outro_dia_nao_reusa_automatico(self):
        admin = criar_usuario('admin_outro_dia', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 2000)
        primeiro = self.novo_processo(self.especie_liq, numero_processo='11g/2026')
        self._salvar_liquidacao(primeiro)
        self.assertEqual(primeiro.numero_relatorio, '2001')
        tramitacao.declinar_analise(primeiro.id, self.analista_liq, 'Teste')
        liberado = LinhaControleRelatorio.objects.get(numero_relatorio='2001')
        liberado.data_relatorio = datetime.date(2020, 1, 1)
        liberado.save(update_fields=['data_relatorio'])

        segundo = self.novo_processo(self.especie_liq, numero_processo='11h/2026')
        self._salvar_liquidacao(segundo)
        self.assertEqual(segundo.numero_relatorio, '2002')
        liberado.refresh_from_db()
        self.assertEqual(liberado.situacao_linha, 'RESERVADA')
        self.assertIsNone(liberado.processo_id)

    def test_numero_antigo_devolvido_nao_vence_o_contador(self):
        admin = criar_usuario('admin_orfao', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 1914)
        # Simula sobra antiga na planilha (como o 1727 que estava voltando).
        LinhaControleRelatorio.objects.create(
            processo=None,
            numero_relatorio='1727',
            numero_processo='orfao/2026',
            grupo='LIQUIDACOES',
            sequencia='LIQUIDACOES',
            data_relatorio=datetime.date(2026, 9, 4),
        )
        processo = self.novo_processo(self.especie_liq, numero_processo='1915x/2026')
        self._salvar_liquidacao(processo)
        self.assertEqual(processo.numero_relatorio, '1915')
        self.assertEqual(relatorios.estado_sequencia()['proximo'], 1916)

    def test_numero_especifico_antigo_nao_puxa_sequencia_para_tras(self):
        admin = criar_usuario('admin_buraco', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 1900)
        primeiro = self.novo_processo(self.especie_liq, numero_processo='93a/2026')
        segundo = self.novo_processo(self.especie_liq, numero_processo='93b/2026')
        self._salvar_liquidacao(primeiro)
        self._salvar_liquidacao(segundo)
        self.assertEqual(primeiro.numero_relatorio, '1901')
        self.assertEqual(segundo.numero_relatorio, '1902')

        # Número específico antigo (fora da sequência atual) não deve
        # fazer o próximo automático voltar a preencher 1727, 1728…
        antigo = self.novo_processo(self.especie_liq, numero_processo='93old/2026')
        tramitacao.assumir(antigo.id, self.analista_liq)
        relatorios.alterar_numero(self.analista_liq, antigo.id, 1726, '2026-09-04')
        antigo.refresh_from_db()
        self.assertEqual(antigo.numero_relatorio, '1726')
        self.assertEqual(relatorios.estado_sequencia()['proximo'], 1903)

        # Número específico à frente só avança o contador; o pulado
        # (1903) não é reaproveitado automaticamente.
        terceiro = self.novo_processo(self.especie_liq, numero_processo='93c/2026')
        tramitacao.assumir(terceiro.id, self.analista_liq)
        relatorios.alterar_numero(self.analista_liq, terceiro.id, 1904, '2026-10-01')
        terceiro.refresh_from_db()
        self.assertEqual(terceiro.numero_relatorio, '1904')
        self.assertEqual(relatorios.estado_sequencia()['proximo'], 1905)

        quarto = self.novo_processo(self.especie_liq, numero_processo='93d/2026')
        self._salvar_liquidacao(quarto)
        self.assertEqual(quarto.numero_relatorio, '1905')

        quinto = self.novo_processo(self.especie_liq, numero_processo='93e/2026')
        self._salvar_liquidacao(quinto)
        self.assertEqual(quinto.numero_relatorio, '1906')

    def test_dois_analistas_nao_recebem_o_mesmo_numero(self):
        outro = criar_usuario('dana_liq', 'ANALISTA_LIQUIDACOES')
        p1 = self.novo_processo(self.especie_liq, numero_processo='14/2026')
        p2 = self.novo_processo(self.especie_liq, numero_processo='15/2026')
        self._salvar_liquidacao(p1, self.analista_liq)
        self._salvar_liquidacao(p2, outro)
        self.assertTrue(p1.numero_relatorio)
        self.assertTrue(p2.numero_relatorio)
        self.assertNotEqual(p1.numero_relatorio, p2.numero_relatorio)

    def test_admin_corrige_numero_de_analise_ja_feita(self):
        admin = criar_usuario('admin_edit', 'GESTAO', is_superuser=True)
        processo = self.novo_processo(self.especie_liq, numero_processo='16/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        relatorios.alterar_numero(admin, processo.id, 3300, '2026-03-15')
        processo.refresh_from_db()
        self.assertEqual(processo.numero_relatorio, '3300')

        self._salvar_liquidacao(processo)
        relatorios.alterar_numero(admin, processo.id, 3300, '2026-03-15')
        processo.refresh_from_db()
        self.assertEqual(processo.numero_relatorio, '3300')
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertEqual(linha.numero_relatorio, '3300')
        self.assertEqual(relatorios.estado_sequencia()['proximo'], 3301)

        with self.assertRaises(PermissionDenied):
            relatorios.alterar_numero(self.analista_lic, processo.id, 3301)

        outro = self.novo_processo(self.especie_liq, numero_processo='17/2026')
        self._salvar_liquidacao(outro)
        relatorios.alterar_numero(admin, outro.id, 3300, '2026-03-16')
        outro.refresh_from_db()
        processo.refresh_from_db()
        self.assertEqual(outro.numero_relatorio, '3300')
        self.assertFalse(processo.numero_relatorio)
        self.assertEqual(
            LinhaControleRelatorio.objects.filter(
                numero_relatorio='3300', situacao_linha='ATIVA').count(), 1)
        self.assertEqual(
            LinhaControleRelatorio.objects.get(numero_relatorio='3300').processo_id,
            outro.id)

    def test_analista_informa_o_proprio_numero(self):
        processo = self.novo_processo(self.especie_liq, numero_processo='16b/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.refresh_from_db()
        relatorios.alterar_numero(self.analista_liq, processo.id, 4400, '2026-04-10')
        processo.refresh_from_db()
        self.assertEqual(processo.numero_relatorio, '4400')
        self.assertEqual(str(processo.data_analise), '2026-04-10')
        self._salvar_liquidacao(processo)
        self.assertEqual(processo.numero_relatorio, '4400')
        relatorios.alterar_numero(self.analista_liq, processo.id, 4401, '2026-04-11')
        processo.refresh_from_db()
        self.assertEqual(processo.numero_relatorio, '4401')
        self.assertEqual(str(processo.data_analise), '2026-04-11')
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertEqual(linha.numero_relatorio, '4401')
        self.assertEqual(str(linha.data_relatorio), '2026-04-11')
        liberado = LinhaControleRelatorio.objects.get(numero_relatorio='4400')
        self.assertEqual(liberado.situacao_linha, 'CANCELADA')
        self.assertIsNone(liberado.processo_id)
        self.assertEqual(str(liberado.data_relatorio), '2026-04-10')

        with self.assertRaises(relatorios.RelatorioInvalido):
            relatorios.alterar_numero(self.analista_liq, processo.id, 4402)

        self.client.force_login(self.analista_liq)
        tela = self.client.get(reverse('analista_processo', args=[processo.id]))
        self.assertContains(tela, 'Informar número específico')
        self.assertContains(tela, 'Data deste número')
        self.assertContains(tela, 'Data do relatório')
        self.assertContains(tela, 'data_relatorio_exibicao')
        self.assertContains(tela, 'btnCancelarNumero')
        self.assertNotContains(tela, 'form="formNumeroRelatorio"')

    def test_salvar_analise_grava_numero_especifico_no_mesmo_botao(self):
        processo = self.novo_processo(self.especie_liq, numero_processo='16c/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        self.client.force_login(self.analista_liq)
        resposta = self.client.post(
            reverse('analista_processo', args=[processo.id]),
            {
                'acao': 'salvar',
                'destino': 'Unidade de Teste',
                'valor': '1000',
                'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
                'numero_relatorio': '5500',
                'data_relatorio': '2026-07-01',
            },
        )
        self.assertEqual(resposta.status_code, 302)
        processo.refresh_from_db()
        self.assertEqual(processo.numero_relatorio, '5500')
        self.assertEqual(str(processo.data_analise), '2026-07-01')

        resposta = self.client.post(
            reverse('analista_processo', args=[processo.id]),
            {
                'acao': 'salvar',
                'destino': 'Unidade de Teste',
                'valor': '1000',
                'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            },
        )
        self.assertEqual(resposta.status_code, 302)
        processo.refresh_from_db()
        self.assertEqual(processo.numero_relatorio, '5500')

    def test_tela_analista_ve_planilha_admin_define_ultimo(self):
        processo = self.novo_processo(self.especie_liq, numero_processo='12/2026')
        self._salvar_liquidacao(processo)

        self.client.force_login(self.analista_liq)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertEqual(pagina.status_code, 200)
        self.assertContains(pagina, 'Planilha — Liquidação')
        self.assertContains(pagina, 'Bolsa Atleta')
        self.assertContains(pagina, 'Inclui (Pagamento Geral; Reanálise).')
        self.assertNotContains(pagina, 'Último nº usado')
        self.assertNotContains(pagina, 'Início da numeração')
        self.assertContains(pagina, 'Números destinados')
        self.assertContains(pagina, 'Análises')
        self.assertNotContains(pagina, 'Do número')
        self.assertContains(pagina, processo.numero_relatorio)
        self.assertContains(pagina, '12/2026')
        self.assertContains(pagina, 'Editar análise')
        self.assertNotContains(pagina, 'name="valor"')
        self.assertNotContains(pagina, 'name="numero_relatorio"')

        admin = criar_usuario('admin_rel2', 'GESTAO', is_superuser=True)
        self.client.force_login(admin)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, 'Numeração')
        self.assertContains(pagina, '12/2026')
        self.assertContains(pagina, 'Editar análise')
        self.assertNotContains(pagina, 'name="numero_relatorio"')
        self.assertNotContains(pagina, 'name="valor"')

        numeracao = self.client.get(
            reverse('controle_relatorio') + '?secao=numeracao')
        self.assertContains(numeracao, 'Início da numeração por grupo')
        self.assertContains(numeracao, 'Último nº usado')
        self.assertContains(numeracao, 'Bolsa Atleta')
        self.assertContains(numeracao, 'Adiantamento')
        self.assertContains(numeracao, 'name="ultimo_numero"')
        resposta = self.client.post(reverse('controle_relatorio_ultimo'), {
            'grupo': 'BOLSA_ATLETA',
            'ultimo_numero': '50',
        })
        self.assertRedirects(
            resposta,
            reverse('controle_relatorio') + '?aba=BOLSA_ATLETA&secao=numeracao')
        self.assertEqual(
            relatorios.estado_sequencia('BOLSA_ATLETA')['proximo'], 51)

    def test_analista_corrige_pela_tela_da_analise(self):
        processo = self.novo_processo(self.especie_liq, numero_processo='18/2026')
        self._salvar_liquidacao(processo)
        tramitacao.liberar_assinatura(processo.id, self.analista_liq)
        processo.refresh_from_db()

        self.client.force_login(self.analista_liq)
        url = reverse('analista_processo', args=[processo.id])
        tela = self.client.get(f'{url}?next={reverse("controle_relatorio")}')
        self.assertEqual(tela.status_code, 200)
        self.assertContains(tela, 'Salvar análise')
        self.assertContains(tela, 'Voltar ao Controle de relatório')

        resposta = self.client.post(url, {
            'acao': 'salvar',
            'next': reverse('controle_relatorio'),
            'destino': 'Secretaria Nova',
            'valor': '2500',
            'periodo': 'Jan/2026',
            'observacao': 'Corrigido na tela da análise',
            'status_analise': 'PROSSEGUIMENTO_COM_RESSALVA',
        })
        self.assertRedirects(resposta, reverse('controle_relatorio'))
        processo.refresh_from_db()
        self.assertEqual(processo.destino, 'Secretaria Nova')
        self.assertEqual(processo.valor, 'R$ 2.500,00')
        self.assertEqual(processo.periodo, 'Jan/2026')
        self.assertEqual(processo.observacao, 'Corrigido na tela da análise')
        self.assertEqual(processo.status_analise, 'PROSSEGUIMENTO_COM_RESSALVA')
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertEqual(linha.destino, 'Secretaria Nova')
        self.assertEqual(linha.valor, 'R$ 2.500,00')
        self.assertEqual(linha.observacao, 'Corrigido na tela da análise')

        self.assertEqual(
            svc_processos.campos_editaveis(self.analista_lic, processo), set())

    def test_destinar_numeros_e_usar_traz_a_data(self):
        p1 = self.novo_processo(self.especie_liq, numero_processo='40/2026')
        p2 = self.novo_processo(self.especie_liq, numero_processo='41/2026')
        p3 = self.novo_processo(self.especie_liq, numero_processo='42/2026')
        relatorios.destinar_numeros(
            self.analista_liq, 5001, 5003, '2026-06-20',
            numeros_processo='40/2026\n41/2026\n42/2026')
        self.assertEqual(list(relatorios.reservas_abertas()), [])
        for p, num in ((p1, '5001'), (p2, '5002'), (p3, '5003')):
            p.refresh_from_db()
            self.assertEqual(p.numero_relatorio, num)
            self.assertEqual(str(p.data_analise), '2026-06-20')
            self.assertTrue(p.sem_relatorio)
            linha = LinhaControleRelatorio.objects.get(processo=p)
            self.assertEqual(linha.numero_relatorio, num)
            self.assertTrue(linha.sem_relatorio)
            self.assertEqual(linha.linha_css_class, 'tr-sem-relatorio')
            self.assertIn('Sem relatório', linha.observacao)

        # Análise preenche dados e mantém amarelo (sem_relatorio).
        tramitacao.assumir(p2.id, self.analista_liq)
        p2.refresh_from_db()
        self.assertEqual(p2.numero_relatorio, '5002')
        svc_processos.aplicar_analise(p2, {
            'destino': 'Unidade de Teste',
            'valor': '1000',
            'status_analise': 'PROSSEGUIMENTO_SEM_RESSALVA',
            'sem_relatorio': '1',
        }, self.analista_liq)
        p2.refresh_from_db()
        self.assertEqual(p2.numero_relatorio, '5002')
        self.assertTrue(p2.sem_relatorio)
        linha2 = LinhaControleRelatorio.objects.get(processo=p2)
        self.assertTrue(linha2.sem_relatorio)
        self.assertEqual(linha2.numero_relatorio, '5002')

        with self.assertRaises(relatorios.RelatorioInvalido) as ctx:
            relatorios.destinar_numeros(
                self.analista_liq, inicio=5001, fim=5001, data='2026-06-21',
                numeros_processo='40/2026')
        self.assertIn('já está em uso', str(ctx.exception))
        with self.assertRaises(PermissionDenied):
            relatorios.destinar_numeros(
                self.analista_lic, inicio=6000, fim=6001, data='2026-06-21',
                numeros_processo='x\ny')

        self.client.force_login(self.analista_liq)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, '5001')
        self.assertContains(pagina, '5002')
        self.assertContains(pagina, 'tr-sem-relatorio')
        self.assertContains(pagina, '40/2026')

        dest = self.client.get(
            reverse('controle_relatorio') + '?aba=destinados')
        self.assertContains(dest, 'Reservar')
        self.assertContains(dest, 'name="processo"')
        self.assertContains(dest, 'Preencher números automaticamente')
        self.assertContains(dest, 'Adicionar processo')
        self.assertContains(dest, 'Grupo de numeração')

        from processos_app.models import EspecieProcesso
        bolsa, _ = EspecieProcesso.objects.get_or_create(
            nome='Concessão Aux. Bolsa Atleta', grupo='LIQUIDACOES',
            defaults={'ativo': True, 'gera_relatorio': True})
        bolsa.gera_relatorio = True
        bolsa.ativo = True
        bolsa.save()
        pb = self.novo_processo(bolsa, numero_processo='bolsa-vinculo/2026')
        relatorios.destinar_numeros(
            self.analista_liq, inicio=11, fim=11, data='2026-08-01',
            grupo='BOLSA_ATLETA', numeros_processo='bolsa-vinculo/2026')
        pb.refresh_from_db()
        self.assertEqual(pb.numero_relatorio, '11')
        aba_bolsa = self.client.get(
            reverse('controle_relatorio') + '?aba=BOLSA_ATLETA')
        self.assertContains(aba_bolsa, '11')
        self.assertContains(aba_bolsa, 'bolsa-vinculo/2026')

    def test_destinar_automatico_usa_cancelados_e_sequencia(self):
        admin = criar_usuario('admin_auto_dest', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 8000)
        # Cancela 8001 → fica vermelho (disponível para auto).
        antigo = self.novo_processo(self.especie_liq, numero_processo='auto-c/2026')
        self._salvar_liquidacao(antigo)
        self.assertEqual(antigo.numero_relatorio, '8001')
        linha = LinhaControleRelatorio.objects.get(processo=antigo)
        relatorios.cancelar_linha(admin, linha.id, 'excluir')
        self.assertNotIn(8001, relatorios.numeros_usados('LIQUIDACOES'))

        p1 = self.novo_processo(self.especie_liq, numero_processo='auto-1/2026')
        p2 = self.novo_processo(self.especie_liq, numero_processo='auto-2/2026')
        resultado = relatorios.destinar_numeros(
            self.analista_liq, data='2026-10-02',
            quantidade_auto=2,
            numeros_processo='auto-1/2026\nauto-2/2026')
        self.assertEqual(resultado['numeros'][0], 8001)
        self.assertEqual(resultado['numeros'][1], 8002)
        p1.refresh_from_db()
        p2.refresh_from_db()
        self.assertEqual(p1.numero_relatorio, '8001')
        self.assertEqual(p2.numero_relatorio, '8002')
        self.assertFalse(
            LinhaControleRelatorio.objects.filter(
                numero_relatorio='8001', situacao_linha='CANCELADA').exists())
        self.assertEqual(
            LinhaControleRelatorio.objects.filter(
                numero_relatorio='8001', sem_relatorio=True).count(), 1)

    def test_especie_define_a_sequencia_do_numero(self):
        """A espécie do cadastro escolhe o contador de relatório."""
        from processos_app.models import EspecieProcesso

        bolsa, _ = EspecieProcesso.objects.get_or_create(
            nome='Concessão Aux. Bolsa Atleta', grupo='LIQUIDACOES',
            defaults={'ativo': True, 'gera_relatorio': True})
        bolsa.gera_relatorio = True
        bolsa.ativo = True
        bolsa.save(update_fields=['gera_relatorio', 'ativo'])
        diaria, _ = EspecieProcesso.objects.get_or_create(
            nome='Concessão Diária', grupo='LIQUIDACOES',
            defaults={'ativo': True, 'gera_relatorio': True})
        diaria.gera_relatorio = True
        diaria.ativo = True
        diaria.save(update_fields=['gera_relatorio', 'ativo'])
        admin = criar_usuario('admin_seq', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 100, grupo='LIQUIDACOES')
        relatorios.definir_ultimo_numero(admin, 10, grupo='BOLSA_ATLETA')
        relatorios.definir_ultimo_numero(admin, 20, grupo='DIARIA')

        liquidacao = self.novo_processo(self.especie_liq, numero_processo='60/2026')
        processo_bolsa = self.novo_processo(bolsa, numero_processo='61/2026')
        processo_diaria = self.novo_processo(diaria, numero_processo='62/2026')
        self._salvar_liquidacao(liquidacao)
        self._salvar_liquidacao(processo_bolsa)
        self._salvar_liquidacao(processo_diaria)

        self.assertEqual(liquidacao.numero_relatorio, '101')
        self.assertEqual(processo_bolsa.numero_relatorio, '11')
        self.assertEqual(processo_diaria.numero_relatorio, '21')
        self.assertEqual(
            LinhaControleRelatorio.objects.get(processo=liquidacao).sequencia,
            'LIQUIDACOES')
        self.assertEqual(
            LinhaControleRelatorio.objects.get(processo=processo_bolsa).sequencia,
            'BOLSA_ATLETA')
        self.assertEqual(
            LinhaControleRelatorio.objects.get(processo=processo_diaria).sequencia,
            'DIARIA')

        self.assertEqual(relatorios.estado_sequencia('LIQUIDACOES')['proximo'], 102)
        self.assertEqual(relatorios.estado_sequencia('BOLSA_ATLETA')['proximo'], 12)
        self.assertEqual(relatorios.estado_sequencia('DIARIA')['proximo'], 22)

        self.client.force_login(self.analista_liq)
        aba_bolsa = self.client.get(
            reverse('controle_relatorio') + '?aba=BOLSA_ATLETA')
        self.assertContains(aba_bolsa, 'Bolsa Atleta')
        self.assertContains(aba_bolsa, '61/2026')
        self.assertContains(aba_bolsa, '>11</strong>')
        self.assertNotContains(aba_bolsa, '60/2026')

        aba_liq = self.client.get(reverse('controle_relatorio'))
        self.assertContains(aba_liq, '60/2026')
        self.assertNotContains(aba_liq, '61/2026')
        self.assertNotContains(aba_liq, '62/2026')

        aba_diaria = self.client.get(
            reverse('controle_relatorio') + '?aba=DIARIA')
        self.assertContains(aba_diaria, '62/2026')
        self.assertContains(aba_diaria, '>21</strong>')

    def test_admin_troca_grupo_de_numeracao_e_gera_novo_numero(self):
        from processos_app.models import EspecieProcesso

        bolsa, _ = EspecieProcesso.objects.get_or_create(
            nome='Concessão Aux. Bolsa Atleta', grupo='LIQUIDACOES',
            defaults={'ativo': True, 'gera_relatorio': True})
        bolsa.gera_relatorio = True
        bolsa.ativo = True
        bolsa.save(update_fields=['gera_relatorio', 'ativo'])
        admin = criar_usuario('admin_grp', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 50, grupo='BOLSA_ATLETA')
        relatorios.definir_ultimo_numero(admin, 80, grupo='SUBVENCAO')

        processo = self.novo_processo(bolsa, numero_processo='70/2026')
        self._salvar_liquidacao(processo)
        self.assertEqual(processo.numero_relatorio, '51')
        self.assertEqual(
            LinhaControleRelatorio.objects.get(processo=processo).sequencia,
            'BOLSA_ATLETA')

        self.client.force_login(admin)
        tela = self.client.get(reverse('analista_processo', args=[processo.id]))
        self.assertContains(tela, 'Trocar grupo de numeração')
        self.assertContains(tela, 'Bolsa Atleta')

        resposta = self.client.post(
            reverse('controle_relatorio_alterar_sequencia', args=[processo.id]), {
                'sequencia': 'SUBVENCAO',
                'next': reverse('analista_processo', args=[processo.id]),
            })
        self.assertEqual(resposta.status_code, 302)
        processo.refresh_from_db()
        self.assertEqual(processo.sequencia_relatorio, 'SUBVENCAO')
        self.assertEqual(processo.numero_relatorio, '81')
        linha = LinhaControleRelatorio.objects.get(processo=processo)
        self.assertEqual(linha.sequencia, 'SUBVENCAO')
        self.assertEqual(linha.numero_relatorio, '81')

        self.client.force_login(self.analista_liq)
        negado = self.client.post(
            reverse('controle_relatorio_alterar_sequencia', args=[processo.id]), {
                'sequencia': 'DIARIA',
            })
        self.assertEqual(negado.status_code, 403)

    def test_planilha_ordena_por_numero_de_relatorio(self):
        admin = criar_usuario('admin_ord', 'GESTAO', is_superuser=True)
        primeiro = self.novo_processo(self.especie_liq, numero_processo='20/2026')
        segundo = self.novo_processo(self.especie_liq, numero_processo='21/2026')
        self._salvar_liquidacao(primeiro)
        self._salvar_liquidacao(segundo)
        relatorios.alterar_numero(admin, primeiro.id, 1895, '2026-05-01')
        relatorios.alterar_numero(admin, segundo.id, 1893, '2026-05-02')
        numeros = [
            linha.numero_relatorio
            for linha in relatorios.listar(self.analista_liq)
            if linha.numero_relatorio in ('1893', '1895')
        ]
        # Mais recentes primeiro (maior número no topo).
        self.assertEqual(numeros, ['1895', '1893'])
        # Números automáticos abandonados ficam vermelhos na planilha.
        self.assertTrue(
            LinhaControleRelatorio.objects.filter(
                situacao_linha='CANCELADA').exclude(
                numero_relatorio__in=('1893', '1895')).exists())

    def test_tela_formata_valor_e_destino_e_lista_de_secretarias(self):
        self.assertEqual(svc_processos.formatar_valor('40000000'), 'R$ 40.000.000,00')
        self.assertEqual(svc_processos.formatar_valor('R$ 40.000.000,00'), 'R$ 40.000.000,00')
        processo = self.novo_processo(self.especie_liq, numero_processo='30/2026')
        tramitacao.assumir(processo.id, self.analista_liq)
        processo.valor = '40000000'
        processo.destino = 'Unidade de Teste'
        processo.save(update_fields=['valor', 'destino'])
        self.client.force_login(self.analista_liq)
        pagina = self.client.get(reverse('analista_processo', args=[processo.id]))
        self.assertContains(pagina, 'R$ 40.000.000,00')
        self.assertContains(pagina, 'name="destino"')
        self.assertContains(pagina, 'Selecione a secretaria')
        self.assertContains(pagina, 'Unidade de Teste')

    def test_cancelar_linha_deixa_vermelho_disponivel(self):
        admin = criar_usuario('admin_cancel', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 5000)
        processo = self.novo_processo(self.especie_liq, numero_processo='c1/2026')
        self._salvar_liquidacao(processo)
        self.assertEqual(processo.numero_relatorio, '5001')

        self.client.force_login(admin)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, 'data-cancelar-linha')
        self.assertContains(pagina, 'Cancelar relatório')
        self.assertNotContains(pagina, 'Guardar número')

        linha = LinhaControleRelatorio.objects.get(processo=processo)
        resp = self.client.post(
            reverse('controle_relatorio_cancelar_linha', args=[linha.id]), {
                'grupo': 'LIQUIDACOES',
            })
        self.assertEqual(resp.status_code, 302)
        linha.refresh_from_db()
        processo.refresh_from_db()
        self.assertEqual(linha.situacao_linha, 'CANCELADA')
        self.assertEqual(linha.numero_relatorio, '5001')
        self.assertTrue(linha.data_relatorio)
        self.assertIsNone(linha.processo_id)
        self.assertFalse(processo.numero_relatorio)
        self.assertNotIn(5001, relatorios.numeros_usados('LIQUIDACOES'))

        planilha = self.client.get(reverse('controle_relatorio'))
        self.assertContains(planilha, 'tr-relatorio-cancelado')
        self.assertContains(planilha, '5001')
        self.assertContains(planilha, 'data-vincular-linha')
        self.assertNotContains(planilha, 'tr-relatorio-reservado')

        novo = self.novo_processo(self.especie_liq, numero_processo='c2/2026')
        tramitacao.assumir(novo.id, self.analista_liq)
        resp2 = self.client.post(
            reverse('controle_relatorio_vincular_linha', args=[linha.id]), {
                'grupo': 'LIQUIDACOES',
                'numero_processo': 'c2/2026',
            })
        self.assertEqual(resp2.status_code, 302)
        linha.refresh_from_db()
        novo.refresh_from_db()
        self.assertEqual(linha.situacao_linha, 'ATIVA')
        self.assertEqual(linha.processo_id, novo.id)
        self.assertEqual(linha.numero_relatorio, '5001')
        self.assertEqual(novo.numero_relatorio, '5001')
        self.assertEqual(linha.numero_processo, 'c2/2026')

    def test_cancelar_linha_historica_libera_numero(self):
        """Linha importada (HISTORICA) também pode ser cancelada."""
        admin = criar_usuario('admin_hist_cancel', 'GESTAO', is_superuser=True)
        linha = LinhaControleRelatorio.objects.create(
            processo=None,
            numero_processo='hist-1/2026',
            numero_relatorio='7201',
            data_relatorio='2026-09-01',
            situacao_linha='HISTORICA',
            sequencia='LIQUIDACOES',
            grupo='LIQUIDACOES',
            objeto='Histórico importado',
        )
        self.assertIn(7201, relatorios.numeros_usados('LIQUIDACOES'))

        self.client.force_login(admin)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, f'data-cancelar-linha="{linha.id}"')

        resp = self.client.post(
            reverse('controle_relatorio_cancelar_linha', args=[linha.id]), {
                'grupo': 'LIQUIDACOES',
            })
        self.assertEqual(resp.status_code, 302)
        linha.refresh_from_db()
        self.assertEqual(linha.situacao_linha, 'CANCELADA')
        self.assertEqual(linha.numero_relatorio, '7201')
        self.assertNotIn(7201, relatorios.numeros_usados('LIQUIDACOES'))

    def test_numero_especifico_reativa_linha_cancelada(self):
        admin = criar_usuario('admin_reativa', 'GESTAO', is_superuser=True)
        relatorios.definir_ultimo_numero(admin, 6100)
        original = self.novo_processo(self.especie_liq, numero_processo='r1/2026')
        self._salvar_liquidacao(original)
        self.assertEqual(original.numero_relatorio, '6101')
        linha = LinhaControleRelatorio.objects.get(processo=original)
        relatorios.cancelar_linha(admin, linha.id, 'excluir')
        linha.refresh_from_db()
        self.assertEqual(linha.situacao_linha, 'CANCELADA')
        self.assertIsNone(linha.processo_id)

        novo = self.novo_processo(self.especie_liq, numero_processo='r2/2026')
        tramitacao.assumir(novo.id, self.analista_liq)
        relatorios.alterar_numero(self.analista_liq, novo.id, 6101, '2026-10-02')
        novo.refresh_from_db()
        linha.refresh_from_db()
        self.assertEqual(novo.numero_relatorio, '6101')
        self.assertEqual(linha.processo_id, novo.id)
        self.assertEqual(linha.situacao_linha, 'ATIVA')
        self.assertEqual(linha.numero_relatorio, '6101')
        self.assertEqual(str(linha.data_relatorio), '2026-10-02')
        self.assertEqual(
            LinhaControleRelatorio.objects.filter(numero_relatorio='6101').count(), 1)
        self.client.force_login(admin)
        planilha = self.client.get(reverse('controle_relatorio'))
        self.assertNotContains(planilha, 'tr-relatorio-cancelado')
        self.assertContains(planilha, '6101')

    def test_admin_apaga_linha_do_controle(self):
        admin = criar_usuario('admin_apaga', 'GESTAO', is_superuser=True)
        gestao = self.gestao
        relatorios.definir_ultimo_numero(admin, 7000)
        processo = self.novo_processo(self.especie_liq, numero_processo='ap1/2026')
        self._salvar_liquidacao(processo)
        self.assertEqual(processo.numero_relatorio, '7001')
        linha = LinhaControleRelatorio.objects.get(processo=processo)

        self.client.force_login(gestao)
        negado = self.client.post(
            reverse('controle_relatorio_apagar_linha', args=[linha.id]), {
                'grupo': 'LIQUIDACOES',
            })
        self.assertEqual(negado.status_code, 302)
        self.assertTrue(LinhaControleRelatorio.objects.filter(id=linha.id).exists())

        self.client.force_login(admin)
        pagina = self.client.get(reverse('controle_relatorio'))
        self.assertContains(pagina, 'data-apagar-linha')
        self.assertContains(pagina, 'Apagar de verdade')

        resp = self.client.post(
            reverse('controle_relatorio_apagar_linha', args=[linha.id]), {
                'grupo': 'LIQUIDACOES',
            })
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(LinhaControleRelatorio.objects.filter(id=linha.id).exists())
        processo.refresh_from_db()
        self.assertFalse(processo.numero_relatorio)
        planilha = self.client.get(reverse('controle_relatorio'))
        self.assertContains(planilha, '0 linhas')
        self.assertContains(planilha, 'Nenhuma linha neste grupo')
        self.assertFalse(
            LinhaControleRelatorio.objects.filter(numero_relatorio='7001').exists())

    def test_protocolo_nao_acessa(self):
        self.client.force_login(self.protocolo)
        self.assertEqual(
            self.client.get(reverse('controle_relatorio')).status_code, 403)

    def test_analista_de_licitacoes_nao_acessa_controle_relatorio(self):
        """Licitações usa Controle de análise; planilha fica com Liquidações."""
        processo = self.novo_processo(self.especie_liq, numero_processo='13/2026')
        self._salvar_liquidacao(processo)
        self.client.force_login(self.analista_lic)
        self.assertEqual(
            self.client.get(reverse('controle_relatorio')).status_code, 403)

    def test_analista_liquidacoes_ve_historico_importado_de_outras_sequencias(self):
        """Importação antiga gravava grupo vazio; analista precisa ver igual ao admin."""
        LinhaControleRelatorio.objects.create(
            processo=None,
            numero_relatorio='72',
            numero_processo='11792/2026',
            contratada='Jade Santos',
            data_relatorio=datetime.date(2026, 7, 1),
            situacao_linha='HISTORICA',
            grupo='',
            sequencia='ADIANTAMENTO',
        )
        LinhaControleRelatorio.objects.create(
            processo=None,
            numero_relatorio='1',
            numero_processo='13024/2025',
            contratada='Laura Atleta',
            data_relatorio=datetime.date(2026, 2, 11),
            situacao_linha='HISTORICA',
            grupo='LIQUIDACOES',
            sequencia='BOLSA_ATLETA',
        )
        self.client.force_login(self.analista_liq)
        adiant = self.client.get(
            reverse('controle_relatorio') + '?aba=ADIANTAMENTO&secao=analises')
        self.assertEqual(adiant.status_code, 200)
        self.assertContains(adiant, '11792/2026')
        self.assertContains(adiant, 'Jade Santos')

        bolsa = self.client.get(
            reverse('controle_relatorio') + '?aba=BOLSA_ATLETA&secao=analises')
        self.assertEqual(bolsa.status_code, 200)
        self.assertContains(bolsa, '13024/2025')
        self.assertContains(bolsa, 'Laura Atleta')

    def test_admin_importa_planilha_excel_historica(self):
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        admin = criar_usuario('admin_import_xlsx', 'GESTAO', is_superuser=True)
        # Linha ativa do sistema não pode ser sobrescrita pela importação.
        relatorios.definir_ultimo_numero(admin, 10)
        vivo = self.novo_processo(self.especie_liq, numero_processo='imp-vivo/2026')
        self._salvar_liquidacao(vivo)
        self.assertEqual(vivo.numero_relatorio, '11')

        wb = Workbook()
        ws = wb.active
        ws.title = 'CONTROLE'
        ws.append(['CONTROLADORIA'])
        ws.append([])
        ws.append([
            'RELATOR', 'Nº relatório', 'DATA', 'PROCESSO ORIGEM',
            'PROCESSO PAGAMENTO', 'SECRETARIA', 'OBJETO', 'PERÍODO',
            'VALOR ', 'DESTINO',
        ])
        ws.append([
            'Ana Teste', 5, datetime.date(2026, 2, 1), '100/2025',
            '200/2026', 'Secretaria X', 'Objeto importado', 'Jan/2026',
            1500.5, 'Destino Y',
        ])
        ws.append([
            'Bruno', 11, datetime.date(2026, 2, 2), '300/2025',
            '400/2026', 'Secretaria Z', 'Não deve sobrescrever', None,
            10, 'Destino Z',
        ])
        ws.append([
            'Carla', '56 A', datetime.date(2026, 2, 3), '500/2025',
            '600/2026', 'Sec W', 'Sufixo letra', None, 99, 'Dest W',
        ])
        # Número pré-preenchido sem dados (como na planilha da CGM) — ignorar.
        ws.append([None, 9999, None, None, None, None, None, None, None, None])
        # Amarelo sem processo = reservado / saiu sem análise — importar.
        from openpyxl.styles import PatternFill
        ws.append([None, 88, datetime.date(2026, 3, 1), None, None, None, None, None, None, None])
        amarelo = PatternFill(
            start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')
        for celula in ws[ws.max_row]:
            celula.fill = amarelo
        buffer = BytesIO()
        wb.save(buffer)
        arquivo = SimpleUploadedFile(
            'controle.xlsx',
            buffer.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'
            ),
        )

        self.client.force_login(admin)
        pagina = self.client.get(reverse('controle_relatorio') + '?secao=numeracao')
        self.assertContains(pagina, 'Importar planilha antiga')
        resp = self.client.post(
            reverse('controle_relatorio_importar'),
            {'grupo': 'LIQUIDACOES', 'planilha': arquivo},
        )
        self.assertEqual(resp.status_code, 302)

        hist = LinhaControleRelatorio.objects.get(numero_relatorio='5')
        self.assertEqual(hist.situacao_linha, 'HISTORICA')
        self.assertEqual(hist.numero_processo, '200/2026')
        self.assertIn('origem: 100/2025', hist.observacao)
        self.assertEqual(hist.analista, 'Ana Teste')
        self.assertTrue(hist.valor.startswith('R$'))

        sufixo = LinhaControleRelatorio.objects.get(numero_relatorio='56 A')
        self.assertEqual(sufixo.situacao_linha, 'HISTORICA')
        self.assertFalse(
            LinhaControleRelatorio.objects.filter(numero_relatorio='9999').exists())
        reservado = LinhaControleRelatorio.objects.get(numero_relatorio='88')
        self.assertEqual(reservado.situacao_linha, 'HISTORICA')
        self.assertTrue(reservado.sem_relatorio)
        self.assertIn('sem análise', reservado.observacao)

        vivo.refresh_from_db()
        self.assertEqual(vivo.numero_relatorio, '11')
        ativo = LinhaControleRelatorio.objects.get(processo=vivo)
        self.assertEqual(ativo.situacao_linha, 'ATIVA')
        self.assertNotEqual(ativo.objeto, 'Não deve sobrescrever')

        # Maior inteiro importado é 88 (amarelo); o 11 ativo já existia.
        self.assertEqual(relatorios.estado_sequencia()['proximo'], 89)
        self.assertIn(5, relatorios.numeros_usados('LIQUIDACOES'))
        self.assertIn(88, relatorios.numeros_usados('LIQUIDACOES'))

        planilha = self.client.get(
            reverse('controle_relatorio') + '?aba=LIQUIDACOES&secao=analises')
        self.assertContains(planilha, 'tr-relatorio-historico')
        self.assertContains(planilha, '200/2026')
        self.assertContains(planilha, '56 A')

    def test_amarelo_na_importacao_so_vale_para_liquidacao(self):
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill

        admin = criar_usuario('admin_amarelo_grp', 'GESTAO', is_superuser=True)
        wb = Workbook()
        ws = wb.active
        ws.title = 'ADIANTAMENTOS'
        ws.append([
            'NUMERO', 'SERVIDOR', ' VALOR ', 'Nº do Processo', 'Assunto',
            'Data', 'Enviado para', 'Secretaria de Origem',
        ])
        ws.append([
            '010', 'Servidor Amarelo', 1000, '99/2026', 'Concessão',
            datetime.date(2026, 1, 6), 'Finanças', 'Sec. Educação',
        ])
        amarelo = PatternFill(
            start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')
        for celula in ws[ws.max_row]:
            celula.fill = amarelo
        buffer = BytesIO()
        wb.save(buffer)
        arquivo = SimpleUploadedFile(
            'adiant_amarelo.xlsx',
            buffer.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'
            ),
        )
        self.client.force_login(admin)
        resp = self.client.post(
            reverse('controle_relatorio_importar'),
            {'grupo': 'ADIANTAMENTO', 'planilha': arquivo},
        )
        self.assertEqual(resp.status_code, 302)
        linha = LinhaControleRelatorio.objects.get(
            sequencia='ADIANTAMENTO', numero_relatorio='010')
        self.assertFalse(linha.sem_relatorio)
        self.assertNotIn('sem análise', (linha.observacao or '').casefold())

    def test_reimporta_substitui_historico_misturado(self):
        """Histórico errado (ex.: Liquidação no Adiantamento) some na reimportação."""
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        admin = criar_usuario('admin_reimp_limpo', 'GESTAO', is_superuser=True)
        LinhaControleRelatorio.objects.create(
            processo=None,
            numero_relatorio='1932',
            numero_processo='13289/2026',
            objeto='aquisição material (não é adiantamento)',
            data_relatorio=datetime.date(2026, 10, 2),
            situacao_linha='HISTORICA',
            grupo='',
            sequencia='ADIANTAMENTO',
        )
        from processos_app.models import SequenciaRelatorio
        SequenciaRelatorio.objects.update_or_create(
            grupo='ADIANTAMENTO', defaults={'proximo_numero': 1933})

        wb = Workbook()
        ws = wb.active
        ws.title = 'ADIANTAMENTOS'
        ws.append([
            'NUMERO', 'SERVIDOR', ' VALOR ', 'Nº do Processo', 'Assunto',
            'Data', 'Enviado para', 'Secretaria de Origem',
        ])
        ws.append([
            '072', 'Jade Santos', 3000, '11792/2026', 'Concessão',
            datetime.date(2026, 7, 1), 'Planejamento', 'Sec. Agricultura',
        ])
        buffer = BytesIO()
        wb.save(buffer)
        arquivo = SimpleUploadedFile(
            'adiant.xlsx',
            buffer.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'
            ),
        )
        self.client.force_login(admin)
        resp = self.client.post(
            reverse('controle_relatorio_importar'),
            {'grupo': 'ADIANTAMENTO', 'planilha': arquivo},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(
            LinhaControleRelatorio.objects.filter(
                sequencia='ADIANTAMENTO', numero_relatorio='1932').exists())
        linha = LinhaControleRelatorio.objects.get(
            sequencia='ADIANTAMENTO', numero_relatorio='072')
        self.assertEqual(linha.contratada, 'Jade Santos')
        self.assertEqual(relatorios.estado_sequencia('ADIANTAMENTO')['proximo'], 73)

    def test_importa_planilha_adiantamento_aba_correta(self):
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        admin = criar_usuario('admin_imp_adiant', 'GESTAO', is_superuser=True)
        wb = Workbook()
        ws_med = wb.active
        ws_med.title = 'MEDIÇÃO ANUAL-ADIANTAMENTO'
        ws_med.append(['NUMERO', 'SERVIDOR', ' VALOR ', 'Nº do Processo', 'Data'])
        ws_med.append(['999', 'Não deve usar esta aba', 1, '1/2026', datetime.date(2026, 1, 1)])
        ws_adi = wb.create_sheet('ADIANTAMENTOS')
        ws_adi.append(['Título qualquer'])
        ws_adi.append([
            'NUMERO', 'SERVIDOR', ' VALOR ', 'Nº do Processo', 'Assunto',
            'Data', 'Enviado para', 'Secretaria de Origem',
        ])
        ws_adi.append([
            '001', 'Servidor Teste', 1000, '21257/2025', 'Concessão',
            datetime.date(2026, 1, 6), 'Finanças', 'Sec. Educação',
        ])
        buffer = BytesIO()
        wb.save(buffer)
        arquivo = SimpleUploadedFile(
            'controle2026.xlsx',
            buffer.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'
            ),
        )
        self.client.force_login(admin)
        resp = self.client.post(
            reverse('controle_relatorio_importar'),
            {'grupo': 'ADIANTAMENTO', 'planilha': arquivo},
        )
        self.assertEqual(resp.status_code, 302)
        linha = LinhaControleRelatorio.objects.get(
            sequencia='ADIANTAMENTO', numero_relatorio='001')
        self.assertEqual(linha.numero_processo, '21257/2025')
        self.assertEqual(linha.contratada, 'Servidor Teste')
        self.assertEqual(linha.grupo, 'LIQUIDACOES')
        self.assertFalse(
            LinhaControleRelatorio.objects.filter(numero_relatorio='999').exists())

        # Analista de Liquidações também vê o histórico importado.
        self.client.force_login(self.analista_liq)
        planilha = self.client.get(
            reverse('controle_relatorio') + '?aba=ADIANTAMENTO&secao=analises')
        self.assertContains(planilha, '21257/2025')
        self.assertContains(planilha, 'Servidor Teste')

    def test_importa_planilha_todas_abas(self):
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        admin = criar_usuario('admin_imp_todas', 'GESTAO', is_superuser=True)
        wb = Workbook()
        ws = wb.active
        ws.title = 'DIÁRIAS'
        ws.append([
            'NUMERO', 'SECRETARIA', 'SERVIDOR', 'QUANT', 'VALOR TOTAL',
            'Nº do Processo', 'Assunto', 'Data', 'Enviado para',
        ])
        ws.append([
            1, 'Sec. Social', 'Maria Silva', 3, 500, '100/2026',
            'Concessão Diária', datetime.date(2026, 3, 1), 'Gabinete',
        ])
        ws2 = wb.create_sheet('BLOCOS')
        ws2.append([
            'NUMERO', 'BLOCO', ' VALOR ', 'Nº do Processo', 'Assunto',
            'Data', 'Enviado para',
        ])
        ws2.append([
            '001', 'Bloco Teste', 7000, '7079/2025', 'PC 2025',
            datetime.date(2026, 2, 3), 'Esporte',
        ])
        ws3 = wb.create_sheet('COTA PATROCINIO')
        ws3.append([
            'NUMERO', 'INTERESSADO', 'EVENTO', ' VALOR ', 'Nº do Processo',
            'Assunto', 'Data', 'Enviado para', 'OBS',
        ])
        ws3.append([
            '001', 'Fed Teste', 'Copa X', 1000, '24589/2025',
            'Prestação de Contas', datetime.date(2026, 1, 9), 'Esporte', 'ok',
        ])
        ws4 = wb.create_sheet('BOLSA ATLETA')
        ws4.append([
            'Número', 'Nome do atleta', 'Nome do responsável', 'VALOR ',
            'Processo de concessão', 'Processo de prestação', 'Assunto',
            'Modalidade', 'Data', 'Enviado para', 'OBS',
        ])
        ws4.append([
            '001', 'Laura Atleta', 'Aline Resp', 3000,
            '22.950/2024', '13024/2025', 'Prestação 2º semestre',
            'Canoa', datetime.date(2026, 2, 11), 'Esporte', '',
        ])
        ws5 = wb.create_sheet('SUBVENÇÃO')
        ws5.append([
            'NUMERO', 'ENTIDADE', ' VALOR ', 'Nº do Processo', 'Assunto',
            'NECESSIDADE CONCESSÃO/PRESTAÇÃO', 'Data', 'Enviado para', 'OBS',
        ])
        ws5.append([
            '001', 'Lar Teste', 197400, '17435/2018', 'Renovação',
            'Renovação para o ano de 2026', datetime.date(2025, 12, 22),
            'PGM', '',
        ])
        buffer = BytesIO()
        wb.save(buffer)
        arquivo = SimpleUploadedFile(
            'multi.xlsx',
            buffer.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'
            ),
        )
        self.client.force_login(admin)
        resp = self.client.post(
            reverse('controle_relatorio_importar'),
            {'grupo': relatorios.GRUPO_IMPORTACAO_COMPLETA, 'planilha': arquivo},
        )
        self.assertEqual(resp.status_code, 302)

        diaria = LinhaControleRelatorio.objects.get(
            sequencia='DIARIA', numero_relatorio='1')
        self.assertEqual(diaria.periodo, '3')
        self.assertEqual(diaria.contratada, 'Maria Silva')

        self.assertTrue(
            LinhaControleRelatorio.objects.filter(
                sequencia='BLOCOS_CARNAVALESCOS', numero_relatorio='001').exists())

        cota = LinhaControleRelatorio.objects.get(
            sequencia='COTA_PATROCINIO', numero_relatorio='001')
        self.assertEqual(cota.objeto, 'Copa X')
        self.assertEqual(cota.periodo, 'Prestação de Contas')
        self.assertEqual(cota.contratada, 'Fed Teste')

        bolsa = LinhaControleRelatorio.objects.get(
            sequencia='BOLSA_ATLETA', numero_relatorio='001')
        self.assertEqual(bolsa.numero_processo, '13024/2025')
        self.assertEqual(bolsa.volume, '22.950/2024')
        self.assertEqual(bolsa.contratada, 'Laura Atleta')
        self.assertEqual(bolsa.periodo, 'Aline Resp')
        self.assertIn('Canoa', bolsa.objeto)

        subv = LinhaControleRelatorio.objects.get(
            sequencia='SUBVENCAO', numero_relatorio='001')
        self.assertEqual(subv.periodo, 'Renovação')
        self.assertEqual(subv.objeto, 'Renovação para o ano de 2026')
        self.assertEqual(subv.contratada, 'Lar Teste')
