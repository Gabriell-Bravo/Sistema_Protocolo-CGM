# processos_app/services/importar_diligencias.py
"""Importa a planilha CONTROLE DE PENDÊNCIAS (.xlsx) para o sistema.

Colunas esperadas (1ª linha):
  Processo origem | Processo Pagamento | Secretaria | Objeto |
  Pendência | Data da pendência | Envio para Simone | Data de Saneamento

Só cria diligências em processos já cadastrados. Não cadastra processo.
Não duplica diligência já existente no mesmo processo com o mesmo texto.
"""

from __future__ import annotations

from datetime import date, datetime, time

from django.db import transaction
from django.utils import timezone

from ..models import Pendencia, Processo
from . import permissions as perm
from .eventos import registrar_evento_pendencia
from .numeros import normalizar_numero
from .tramitacao import TransicaoInvalida

ESPECIE_STUB = 'Importação de pendências'

CABECALHOS = {
    'processo origem': 'origem',
    'processo pagamento': 'pagamento',
    'secretaria': 'secretaria',
    'objeto': 'objeto',
    'pendência': 'pendencia',
    'pendencia': 'pendencia',
    'data da pendência': 'data_pendencia',
    'data da pendencia': 'data_pendencia',
    'envio para simone': 'envio',
    'data de saneamento': 'data_saneamento',
}


def normalizar_descricao(texto):
    return ' '.join((texto or '').split()).casefold()


def _celula_texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor == int(valor):
        return str(int(valor))
    return str(valor).strip()


def _celula_data(valor):
    if valor is None or valor == '':
        return None
    if isinstance(valor, datetime):
        return timezone.make_aware(
            datetime.combine(valor.date(), time.min)
        ) if timezone.is_naive(valor) else valor
    if isinstance(valor, date):
        return timezone.make_aware(datetime.combine(valor, time.min))
    return None


def _mapear_cabecalho(linha):
    mapa = {}
    for indice, bruto in enumerate(linha or []):
        chave = CABECALHOS.get(_celula_texto(bruto).casefold())
        if chave and chave not in mapa:
            mapa[chave] = indice
    return mapa


def _ler_linhas(arquivo):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise TransicaoInvalida(
            'Biblioteca openpyxl indisponível no servidor.') from exc

    try:
        workbook = load_workbook(arquivo, data_only=True, read_only=True)
    except Exception as exc:
        raise TransicaoInvalida(
            'Não foi possível ler o arquivo Excel.') from exc

    planilha = workbook.active
    linhas = list(planilha.iter_rows(values_only=True))
    if not linhas:
        raise TransicaoInvalida('Planilha vazia.')

    mapa = _mapear_cabecalho(linhas[0])
    if 'pendencia' not in mapa or (
            'origem' not in mapa and 'pagamento' not in mapa):
        raise TransicaoInvalida(
            'Cabeçalho inválido. Espere colunas Processo origem / '
            'Processo Pagamento e Pendência.')

    registros = []
    for bruto in linhas[1:]:
        if not bruto:
            continue

        def col(nome):
            indice = mapa.get(nome)
            if indice is None or indice >= len(bruto):
                return None
            return bruto[indice]

        origem = normalizar_numero(col('origem'))
        pagamento = normalizar_numero(col('pagamento'))
        descricao = _celula_texto(col('pendencia'))
        if not descricao or (not origem and not pagamento):
            continue
        registros.append({
            'origem': origem,
            'pagamento': pagamento,
            'secretaria': _celula_texto(col('secretaria')) or 'Não informada',
            'objeto': _celula_texto(col('objeto')) or 'Importado da planilha de pendências',
            'descricao': descricao,
            'data_pendencia': _celula_data(col('data_pendencia')),
            'data_saneamento': _celula_data(col('data_saneamento')),
        })
    return registros


def _indice_processos():
    """numero normalizado → Processo (pagamento e origem). Ignora stubs."""
    indice = {}
    for processo in (
            Processo.objects.exclude(especie=ESPECIE_STUB)
            .only('id', 'numero_processo', 'processo_pagamento')):
        for bruto in (processo.numero_processo, processo.processo_pagamento):
            chave = normalizar_numero(bruto)
            if chave and chave not in indice:
                indice[chave] = processo
    return indice


def _localizar_processo(registro, indice):
    for chave in (registro['pagamento'], registro['origem']):
        if chave and chave in indice:
            return indice[chave]
    return None


def _descricoes_existentes(processo_ids):
    mapa = {}
    if not processo_ids:
        return mapa
    for pendencia in (
            Pendencia.objects.filter(processo_id__in=processo_ids)
            .exclude(status='CANCELADA')
            .only('processo_id', 'descricao')):
        mapa.setdefault(pendencia.processo_id, set()).add(
            normalizar_descricao(pendencia.descricao))
    return mapa


def stubs_de_importacao():
    """Processos criados por engano na importação antiga."""
    return Processo.objects.filter(especie=ESPECIE_STUB)


@transaction.atomic
def remover_stubs_de_importacao():
    """Apaga processos stub e as diligências ligadas a eles."""
    consulta = stubs_de_importacao()
    removidos = consulta.count()
    consulta.delete()
    return removidos


@transaction.atomic
def importar_planilha(usuario, arquivo):
    """Importa só pendências em processos já existentes. Retorna contagens."""
    perm.assert_permissao(
        perm.pode_concluir_pendencia(usuario),
        'Somente a Gestão ou o administrador importam a planilha de diligências.')

    if not arquivo:
        raise TransicaoInvalida('Selecione um arquivo .xlsx.')

    registros = _ler_linhas(arquivo)
    if not registros:
        raise TransicaoInvalida('Nenhuma linha de pendência encontrada na planilha.')

    indice = _indice_processos()
    criadas = 0
    ignoradas = 0
    sem_processo = 0
    saneadas = 0
    vistos_arquivo = set()

    resolvidos = []
    for registro in registros:
        processo = _localizar_processo(registro, indice)
        if processo is None:
            sem_processo += 1
            continue
        resolvidos.append((registro, processo))

    existentes = _descricoes_existentes({p.id for _, p in resolvidos})

    for registro, processo in resolvidos:
        chave_dup = (processo.id, normalizar_descricao(registro['descricao']))
        if chave_dup in vistos_arquivo:
            ignoradas += 1
            continue
        vistos_arquivo.add(chave_dup)

        if chave_dup[1] in existentes.get(processo.id, set()):
            ignoradas += 1
            continue

        pendencia = Pendencia(
            processo=processo,
            descricao=registro['descricao'],
            criada_por=usuario,
            responsavel_tecnico=None,
            status='AGUARDANDO_ATENDIMENTO',
        )
        if registro['data_saneamento']:
            pendencia.status = 'RESOLVIDA'
            pendencia.resolvida_em = registro['data_saneamento']
            pendencia.resolvida_por = usuario
            saneadas += 1

        pendencia.save()
        if registro['data_pendencia']:
            Pendencia.objects.filter(id=pendencia.id).update(
                criada_em=registro['data_pendencia'])

        registrar_evento_pendencia(
            pendencia, 'CRIADA', usuario, 'Importada da planilha.')
        if registro['data_saneamento']:
            registrar_evento_pendencia(
                pendencia, 'RESOLVIDA', usuario,
                f'Concluída por {perm.nome_usuario(usuario)}.')
        existentes.setdefault(processo.id, set()).add(chave_dup[1])
        criadas += 1

    return {
        'criadas': criadas,
        'ignoradas': ignoradas,
        'sem_processo': sem_processo,
        'saneadas': saneadas,
        'lidas': len(registros),
    }
