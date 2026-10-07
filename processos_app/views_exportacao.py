# processos_app/views_exportacao.py
"""Tela HTML de exportação/impressão de saídas + Excel alinhado."""

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .services import exportacao_saidas as svc


def _assert_pode_exportar(user):
    if not svc.pode_exportar(user):
        raise PermissionDenied(
            'Somente Protocolo, Gestão e o administrador exportam saídas.')


@login_required
@require_GET
def exportar_saidas(request):
    """Hub com modelos prontos + preview A4 sempre visível."""
    _assert_pode_exportar(request.user)
    try:
        filtros = svc.filtros_de_request(request)
    except svc.ExportacaoInvalida as exc:
        # Período personalizado sem datas: mostra tela com aviso.
        filtros = {
            'periodo': 'personalizado',
            'genero': (request.GET.get('genero') or '').strip(),
            'sequencia': '',
            'especie': (request.GET.get('especie') or '').strip(),
            'termo': (request.GET.get('termo') or '').strip(),
            'data_inicial': (request.GET.get('data_inicial') or '').strip(),
            'data_final': (request.GET.get('data_final') or '').strip(),
            'intervalo': (None, None),
        }
        if filtros['genero'] != 'LIQUIDACOES':
            filtros['sequencia'] = ''
        else:
            filtros['sequencia'] = (request.GET.get('sequencia') or '').strip()
        processos = []
        erro = str(exc)
        total = 0
    else:
        erro = ''
        qs = svc.consultar_saidas(request.user, filtros)
        processos = list(qs)
        total = len(processos)

    qs_base = svc.querystring_exportacao(filtros)
    return render(request, 'finalizados_exportar.html', {
        'filtros': filtros,
        'processos': processos,
        'total': total,
        'erro': erro,
        'titulo': svc.titulo_relatorio(filtros),
        'rotulo_periodo': svc.rotulo_periodo(filtros),
        'periodos': svc.periodos_disponiveis(),
        'grupos': svc.grupos_disponiveis(request.user),
        'subgrupos': (
            svc.subgrupos_liquidacoes()
            if filtros.get('genero') == 'LIQUIDACOES' else []
        ),
        'especies': svc.especies_do_conjunto(request.user, filtros),
        'emitido_em': timezone.localtime(),
        'querystring': qs_base,
        'url_excel': f'{reverse("exportar_finalizados_excel")}?{qs_base}',
    })


@login_required
@require_GET
def exportar_finalizados_excel(request):
    """Excel com os mesmos filtros da tela de impressão."""
    _assert_pode_exportar(request.user)
    try:
        filtros = svc.filtros_de_request(request)
    except svc.ExportacaoInvalida as exc:
        return HttpResponse(
            f'{{"success": false, "message": "{exc}"}}',
            content_type='application/json', status=400)

    processes = svc.consultar_saidas(request.user, filtros)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Processos Finalizados'

    headers = [
        'N° Processo', 'Volume', 'Secretaria', 'Data Entrada', 'Hora Entrada',
        'Data Saída', 'Hora Saída', 'Destino', 'Gênero', 'Espécie', 'Objeto',
        'Contratada', 'Recorrente', 'Prioridade', 'Técnico', 'N° Despacho',
        'Observação', 'Valor',
    ]
    sheet.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(
        start_color='4CAF50', end_color='4CAF50', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    for col_num, _header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for process in processes:
        sheet.append([
            process.numero_processo,
            process.volume,
            process.secretaria,
            process.data_entrada.strftime('%Y-%m-%d') if process.data_entrada else '',
            process.hora_entrada.strftime('%H:%M') if process.hora_entrada else '',
            process.data_saida.strftime('%Y-%m-%d') if process.data_saida else '',
            process.hora_saida.strftime('%H:%M') if process.hora_saida else '',
            process.destino,
            process.genero,
            process.especie,
            process.objeto,
            process.contratada,
            process.recorrente,
            process.prioridade,
            process.tecnico or process.nome_analista,
            process.numero_despacho,
            process.observacao_protocolo or process.observacao or '',
            process.valor,
        ])

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border

    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        sheet.column_dimensions[column_letter].width = min(max_length + 2, 100)

    ini, fim = filtros.get('intervalo') or (None, None)
    if ini and fim:
        nome = f'processos_finalizados_{ini.isoformat()}_a_{fim.isoformat()}.xlsx'
    else:
        nome = 'processos_finalizados_todas.xlsx'

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        ),
    )
    response['Content-Disposition'] = f'attachment; filename={nome}'
    workbook.save(response)
    return response
